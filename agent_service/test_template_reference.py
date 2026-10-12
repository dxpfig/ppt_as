"""Company template references stay bounded and never become the deck."""
from pathlib import Path
import json
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt

from agent_service import template_reference as references


def sample_template(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    heading = cover.shapes.add_textbox(Pt(60), Pt(200), Pt(800), Pt(90))
    run = heading.text_frame.paragraphs[0].add_run()
    run.text = '季度经营汇报'
    run.font.size = Pt(40)
    run.font.name = '微软雅黑'
    run.font.color.rgb = RGBColor.from_string('E50011')
    body = prs.slides.add_slide(prs.slide_layouts[6])
    title = body.shapes.add_textbox(Pt(60), Pt(40), Pt(400), Pt(40))
    title.text_frame.paragraphs[0].add_run().text = '第一章 核心结论'
    card = body.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Pt(60), Pt(120), Pt(260), Pt(160))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string('F5F7FA')
    text = body.shapes.add_textbox(Pt(80), Pt(140), Pt(220), Pt(120))
    text.text_frame.paragraphs[0].add_run().text = '收入同比增长 12%，主要来自存量客户续约。'
    closing = prs.slides.add_slide(prs.slide_layouts[6])
    end = closing.shapes.add_textbox(Pt(60), Pt(220), Pt(800), Pt(60))
    end.text_frame.paragraphs[0].add_run().text = '下一步行动'
    prs.save(path)
    return path


class TemplateReferenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = sample_template(self.root / 'template.pptx')

    def tearDown(self):
        self.tmp.cleanup()

    def test_digest_is_bounded_brand_aware_and_indexed(self):
        references.build(self.source, self.root, 'corp')
        reference = references.load(self.root, 'corp')
        summary = reference['summary']
        self.assertEqual(summary['kind'], 'company-template-reference')
        self.assertEqual(summary['pageCount'], 3)
        self.assertEqual(summary['pages'][0]['sourceId'], 'template-page-01')
        self.assertIn('#E50011', [item['color'] for item in summary['palette']])
        self.assertIn('微软雅黑', [item['family'] for item in summary['fonts']])
        self.assertEqual([item['archetype'] for item in summary['pages']], ['cover', 'content', 'closing'])
        self.assertEqual(len(reference['sources']), 3)
        # The digest carries page purpose, never the template's body copy.
        self.assertIn('季度经营汇报', summary['pages'][0]['title'])
        digest = json.dumps(summary, ensure_ascii=False)
        self.assertNotIn('存量客户续约', digest)
        self.assertLess(len(digest), 4000)

    def test_pages_remain_readable_on_demand(self):
        references.build(self.source, self.root, 'corp')
        reference = references.load(self.root, 'corp')
        page = json.loads(reference['sources']['template-page-02'])
        texts = [element['content']['text'] for element in page['elements'] if element['elementType'] == 'text']
        self.assertTrue(any('存量客户续约' in text for text in texts))

    def test_build_is_idempotent_for_the_same_template(self):
        first = references.build(self.source, self.root, 'corp')
        second = references.build(self.source, self.root, 'corp')
        self.assertEqual(first['pageCount'], second['pageCount'])

    def test_missing_reference_is_not_invented(self):
        self.assertIsNone(references.load(self.root, 'unknown'))
        self.assertIsNone(references.load(None, 'corp'))

    def test_a_reference_from_an_older_importer_is_measured_again(self):
        references.build(self.source, self.root, 'corp')
        self.assertFalse(references.stale(self.root, 'corp'))
        stamp = self.root / references.REFERENCE_DIRECTORY / 'corp' / 'reference.json'
        entry = json.loads(stamp.read_text(encoding='utf-8'))
        entry['format'] = 'pptd-importer-1'
        stamp.write_text(json.dumps(entry, ensure_ascii=False), encoding='utf-8')
        self.assertTrue(references.stale(self.root, 'corp'))
        # Rebuilding restores the current stamp and keeps the template usable.
        references.build(self.source, self.root, 'corp')
        self.assertFalse(references.stale(self.root, 'corp'))
        self.assertIsNotNone(references.load(self.root, 'corp'))


if __name__ == '__main__':
    unittest.main()
