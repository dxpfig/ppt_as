"""Image pipeline checks without contacting external providers or user knowledge bases."""
import asyncio
import html
from io import BytesIO
import json
import os
import socket
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

import httpx
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from agent_service import images
from agent_service.models import Project, Slide, SlideImage
from agent_service.rendering import create_deck


def bitmap():
    output = BytesIO()
    Image.new("RGB", (800, 500), "#6b967b").save(output, format="PNG")
    return output.getvalue()


class ImageTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.env = patch.dict(os.environ, {"WEKNORA_BASE_URL": "http://knowledge.test", "WEKNORA_API_KEY": "test-secret",
                                         "WEKNORA_TENANT_ID": "42", "WEKNORA_KNOWLEDGE_BASE_IDS": "kb-a,kb-b",
                                         "PPT_IMAGE_WEB_PROVIDERS": "mcp,bing", "PPT_WEB_MCP_URL": "http://search.test/mcp"})
        self.env.start()

    async def asyncTearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    async def test_knowledge_results_include_image_info_and_resource_handles(self):
        payload = {"data": [{"knowledge_title": "架构", "image_info": json.dumps([{"url": "resource://architecture", "caption": "系统架构"}]),
                              "content": "![流程](resource://flow)"}]}
        with patch.object(images, "api_json", new=AsyncMock(return_value=payload)) as request:
            found = await images.knowledge_search("系统架构")
        self.assertEqual(len(found), 2)
        self.assertIn("file_path=resource%3A%2F%2Farchitecture", found[0].url)
        self.assertTrue(found[0].private)
        self.assertEqual(request.call_args.kwargs["headers"]["X-Tenant-ID"], "42")
        self.assertEqual(request.call_args.kwargs["json"]["knowledge_base_ids"], ["kb-a", "kb-b"])

    async def test_cache_validates_bitmap_and_does_not_expose_remote_credentials(self):
        with patch.object(images, "fetch_bytes", new=AsyncMock(return_value=bitmap())):
            result = await images.cache_candidate(self.root, images.Candidate("http://knowledge.test/api/v1/files?file_path=resource://private", "架构图", "weknora", private=True))
        self.assertEqual((result.width, result.height), (800, 500))
        self.assertTrue(images.asset_path(self.root, result.assetId).exists())
        self.assertEqual(result.sourceUrl, "")
        self.assertNotIn("test-secret", result.model_dump_json())
        self.assertTrue(result.url.startswith("/api/studio/images/"))
        with self.assertRaises(Exception):
            images.normalize_bitmap(b"<html>not an image</html>")
        with self.assertRaises(images.ImageSearchError):
            images.asset_path(self.root, "../../etc/passwd")

    async def test_bing_skips_stock_previews_and_reaches_later_usable_results(self):
        rows = [{"murl": f"https://www.vecteezy.com/preview-{i}.jpg", "purl": "https://www.vecteezy.com/photo"} for i in range(20)]
        rows.append({"murl": "https://photos.test/solar.jpg", "purl": "https://source.test/solar", "t": "Rooftop solar"})
        markup = "<html>" + "".join(f'<a class="iusc" m="{html.escape(json.dumps(row), quote=True)}"></a>' for row in rows) + "</html>"
        with patch.object(images, "fetch_bytes", new=AsyncMock(return_value=markup.encode())):
            found = await images.bing_search("solar")
        self.assertEqual([item.url for item in found], ["https://photos.test/solar.jpg"])
        self.assertEqual(found[0].source_url, "https://source.test/solar")

    async def test_bing_rejects_unrelated_search_results(self):
        rows = [{"murl": "https://img.test/actress.jpg", "purl": "https://fashion.test/hair", "t": "明星造型与时尚穿搭"},
                {"murl": "https://img.test/solar.jpg", "purl": "https://energy.test/project", "t": "工业园屋顶光伏建设"}]
        markup = "<html>" + "".join(f'<a class="iusc" m="{html.escape(json.dumps(row), quote=True)}"></a>' for row in rows) + "</html>"
        with patch.object(images, "fetch_bytes", new=AsyncMock(return_value=markup.encode())):
            found = await images.bing_search("industrial park rooftop solar installation aerial view")
        self.assertEqual([item.url for item in found], ["https://img.test/solar.jpg"])

    async def test_credentials_are_not_forwarded_on_external_redirect(self):
        requests = []
        def handle(request):
            requests.append(request)
            if request.url.host == "knowledge.test":
                return httpx.Response(302, headers={"Location": "https://cdn.test/photo.jpg"})
            return httpx.Response(200, content=bitmap())
        client = httpx.AsyncClient(transport=httpx.MockTransport(handle))
        with patch.object(images.httpx, "AsyncClient", return_value=client), patch.object(images, "public_url", new=AsyncMock()) as validate:
            await images.fetch_bytes("http://knowledge.test/api/v1/files?file_path=resource://a", private=True)
        self.assertEqual(requests[0].headers["X-API-Key"], "test-secret")
        self.assertNotIn("X-API-Key", requests[1].headers)
        validate.assert_awaited_once_with("https://cdn.test/photo.jpg")

    async def test_private_network_image_urls_are_rejected(self):
        addresses = [(2, 1, 6, "", ("127.0.0.1", 80))]
        with patch.object(images.socket, "getaddrinfo", return_value=addresses), patch.dict(os.environ, {"PPT_IMAGE_ALLOWED_HOSTS": ""}):
            with self.assertRaises(images.ImageSearchError):
                await images.public_url("http://untrusted.test/image.png")

    async def test_mcp_precedes_fallback_and_knowledge_only_never_uses_web(self):
        calls = []
        async def provider(name, query):
            calls.append(name)
            if name in ("weknora", "mcp"):
                raise images.ImageSearchError("检索暂不可用")
            return [images.Candidate("https://photos.test/solar.jpg", "太阳能", name)]
        with patch.object(images, "provider_search", side_effect=provider), patch.object(images, "fetch_bytes", new=AsyncMock(return_value=bitmap())):
            found, _ = await images.search_images(self.root, "solar panels", "auto", 1)
            self.assertEqual(calls, ["weknora", "mcp", "bing"])
            self.assertEqual(found[0].provider, "bing")
            calls.clear()
            found, _ = await images.search_images(self.root, "内部图", "knowledge", 1)
            self.assertEqual(calls, ["weknora"])
            self.assertEqual(found, [])

    async def test_http_mcp_discovery_and_image_search_protocol(self):
        from mcp.server.fastmcp import FastMCP
        import uvicorn
        mcp = FastMCP("image-search-test", stateless_http=True, json_response=True)

        @mcp.tool()
        def image_search(query: str) -> dict:
            return {"images": [{"url": "https://photos.test/solar.jpg", "title": query}]}

        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(mcp.streamable_http_app(), log_level="error", lifespan="on"))
        task = asyncio.create_task(server.serve(sockets=[sock]))
        try:
            async with asyncio.timeout(5):
                while not server.started:
                    await asyncio.sleep(.01)
            with patch.dict(os.environ, {"PPT_WEB_MCP_URL": f"http://127.0.0.1:{port}/mcp", "PPT_WEB_MCP_HEADERS": "{}", "PPT_WEB_MCP_ARGS": "{}", "PPT_WEB_MCP_TRANSPORT": "http", "PPT_WEB_MCP_TOOL": "", "PPT_WEB_MCP_QUERY_KEY": ""}):
                found = await images.mcp_search("solar panels")
            self.assertEqual(found[0].url, "https://photos.test/solar.jpg")
            self.assertEqual(found[0].title, "solar panels")
        finally:
            server.should_exit = True
            await task
            sock.close()

    async def test_optional_search_adapters_normalize_results_and_send_correct_auth(self):
        with patch.dict(os.environ, {"TAVILY_API_KEY": "tavily-test", "SERPER_API_KEY": "serper-test", "SEARXNG_BASE_URL": "http://search.test"}):
            with patch.object(images, "api_json", new=AsyncMock(return_value={"images": [{"url": "https://img.test/t.jpg", "description": "Solar"}]})) as api:
                found = await images.provider_search("tavily", "solar")
                self.assertEqual(found[0].url, "https://img.test/t.jpg")
                self.assertTrue(api.call_args.kwargs["json"]["include_images"])
                self.assertEqual(api.call_args.kwargs["headers"]["Authorization"], "Bearer tavily-test")
            with patch.object(images, "api_json", new=AsyncMock(return_value={"images": [{"imageUrl": "https://img.test/s.jpg", "link": "https://page.test/s"}]})) as api:
                found = await images.provider_search("serper", "solar")
                self.assertEqual(found[0].source_url, "https://page.test/s")
                self.assertEqual(api.call_args.kwargs["headers"]["X-API-KEY"], "serper-test")
            with patch.object(images, "api_json", new=AsyncMock(return_value={"results": [{"img_src": "https://img.test/x.jpg", "url": "https://page.test/x"}]})) as api:
                found = await images.provider_search("searxng", "solar")
                self.assertEqual(found[0].url, "https://img.test/x.jpg")
                self.assertEqual(found[0].source_url, "https://page.test/x")
                self.assertEqual(api.call_args.kwargs["params"]["categories"], "images")

    async def test_picture_is_embedded_and_survives_content_export(self):
        with patch.object(images, "fetch_bytes", new=AsyncMock(return_value=bitmap())):
            image = await images.cache_candidate(self.root, images.Candidate("https://photos.test/solar.jpg", "光伏组件", "commons", "https://source.test/photo", "Author", "CC BY 4.0"))
        slide = Slide(id="page", title="清洁能源", points=["太阳能发电"], image=image)
        project = Project(id="image-check", slides=[slide])
        deck = self.root / "illustrated.pptx"
        create_deck(project, deck, None, {slide.id: images.asset_path(self.root, image.assetId)})
        prs = Presentation(deck)
        self.assertEqual(sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in prs.slides[0].shapes), 1)
        self.assertIn("CC BY 4.0", prs.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn("https://source.test/photo", prs.slides[0].notes_slide.notes_text_frame.text)

    async def test_removed_or_disabled_images_do_not_reappear(self):
        project = Project(id="skip", slides=[Slide(id="one", title="无需配图", imageDisabled=True)])
        with patch.object(images, "search_images", new=AsyncMock()) as search:
            _, pictures, _ = await images.illustrate_project(self.root, project, lambda *args: None)
            self.assertEqual(pictures, {})
            search.assert_not_awaited()
            project.imageSource = "none"
            await images.illustrate_project(self.root, project, lambda *args: None)
            search.assert_not_awaited()
            project.imageSource = "auto"
            project.slides[0].imageDisabled = False
            project.slides[0].imageIntent = "none"
            await images.illustrate_project(self.root, project, lambda *args: None)
            search.assert_not_awaited()

    async def test_unavailable_images_do_not_block_deck_creation(self):
        project = Project(id="no-image", slides=[Slide(id="one", title="主题")])
        with patch.object(images, "search_images", new=AsyncMock(return_value=([], ["服务暂不可用"]))):
            illustrated, pictures, warnings = await images.illustrate_project(self.root, project, lambda *args: None)
        self.assertFalse(pictures)
        self.assertTrue(any("文字版式" in warning for warning in warnings))
        create_deck(illustrated, self.root / "text.pptx", None, pictures)
        self.assertTrue((self.root / "text.pptx").exists())


if __name__ == "__main__":
    unittest.main()
