"""Surgical PPTX edits keep every unnamed part byte-identical."""
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Pt

from agent_service import surgical


def sample(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    first = prs.slides.add_slide(prs.slide_layouts[6])
    first.shapes.add_textbox(Pt(60), Pt(200), Pt(800), Pt(80)).text_frame.text = '关于AI服务器的申购汇报'
    first.shapes.add_textbox(Pt(60), Pt(300), Pt(400), Pt(40)).text_frame.text = 'AI项目组'
    second = prs.slides.add_slide(prs.slide_layouts[6])
    second.shapes.add_textbox(Pt(60), Pt(60), Pt(800), Pt(60)).text_frame.text = '目录'
    extra = second.shapes.add_textbox(Pt(60), Pt(160), Pt(800), Pt(60))
    extra.text_frame.text = '必要性分析'
    prs.save(path)
    return path


class SurgicalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = sample(self.root / 'deck.pptx')

    def tearDown(self):
        self.tmp.cleanup()

    def test_inspect_lists_stable_paragraph_addresses(self):
        report = surgical.inspect(self.source)
        self.assertEqual(len(report['slides']), 2)
        identifiers = [item['id'] for page in report['slides'] for item in page['paragraphs']]
        self.assertIn('s1#p0', identifiers)
        self.assertIn('s2#p1', identifiers)

    def test_named_paragraphs_change_and_everything_else_stays_identical(self):
        before = ZipFile(self.source)
        destination = self.root / 'edited.pptx'
        report = surgical.apply(self.source, [
            {'id': 's1#p0', 'text': '2026 年度 AI 服务器申购汇报'},
            {'slide': 2, 'find': '必要性分析', 'replace': '申购必要性'},
        ], destination)
        self.assertTrue(report['verification']['untouchedPartsIdentical'])
        after = ZipFile(destination)
        self.assertEqual(sorted(before.namelist()), sorted(after.namelist()))
        # Only the two edited slide parts may differ, byte for byte.
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        self.assertEqual(sorted(changed), ['ppt/slides/slide1.xml', 'ppt/slides/slide2.xml'])
        self.assertIn('2026 年度 AI 服务器申购汇报', report['edit'][1])
        # The source file itself is never rewritten.
        self.assertEqual(ZipFile(self.source).read('ppt/slides/slide1.xml'), before.read('ppt/slides/slide1.xml'))

    def test_invalid_operations_are_refused_without_writing(self):
        for operations in ([], [{'id': 's1#p0'}], [{'slide': 0, 'find': 'a', 'replace': 'b'}],
                           [{'id': 's1#p0', 'text': 'x', 'extra': 1}], [{'slide': 1, 'find': '', 'replace': 'b'}]):
            with self.subTest(operations=operations):
                with self.assertRaises(surgical.SurgicalError):
                    surgical.apply(self.source, operations, self.root / 'never.pptx')
        self.assertFalse((self.root / 'never.pptx').exists())

    def test_missing_find_target_is_reported_not_silently_skipped(self):
        with self.assertRaises(surgical.SurgicalError):
            surgical.apply(self.source, [{'slide': 2, 'find': '不存在的句子', 'replace': '新文字'}],
                           self.root / 'never.pptx')
        self.assertFalse((self.root / 'never.pptx').exists())

    def test_source_and_destination_must_differ(self):
        with self.assertRaisesRegex(surgical.SurgicalError, '另一份文件'):
            surgical.apply(self.source, [{'id': 's1#p0', 'text': 'x'}], self.source)

    def test_select_keeps_only_the_chosen_pages_and_their_design(self):
        deck = self.root / 'four.pptx'
        prs = Presentation()
        for index in range(4):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Pt(60), Pt(60), Pt(800), Pt(60)).text_frame.text = f'第 {index + 1} 页'
        prs.save(deck)
        destination = self.root / 'two.pptx'
        report = surgical.select(deck, [1, 3], destination)
        self.assertEqual(report['pages'], [1, 3])
        with ZipFile(deck) as before, ZipFile(destination) as after:
            self.assertEqual(len(surgical.inspect(destination)['slides']), 2)
            # Kept slides are byte-identical; the master/theme relationships stay.
            for name in ('ppt/slides/slide1.xml', 'ppt/slides/slide3.xml', 'ppt/slideMasters/slideMaster1.xml'):
                self.assertEqual(before.read(name), after.read(name), name)
            self.assertNotIn('ppt/slides/slide2.xml', after.namelist())
            self.assertNotIn('ppt/slides/slide4.xml', after.namelist())
            rels = after.read('ppt/_rels/presentation.xml.rels').decode()
            self.assertIn('slideMaster1.xml', rels)
            self.assertNotIn('slide2.xml', rels)
            self.assertEqual(before.namelist(), before.namelist())

    def test_select_reorders_and_refuses_bad_requests(self):
        deck = self.root / 'three.pptx'
        prs = Presentation()
        for index in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Pt(60), Pt(60), Pt(800), Pt(60)).text_frame.text = f'P{index + 1}'
        prs.save(deck)
        destination = self.root / 'reordered.pptx'
        surgical.select(deck, [3, 1], destination)
        titles = [page['paragraphs'][0]['text'] for page in surgical.inspect(destination)['slides']]
        self.assertEqual(titles, ['P3', 'P1'])
        for pages in ([], [0], [1, 1], [9], ['1']):
            with self.subTest(pages=pages):
                with self.assertRaises(surgical.SurgicalError):
                    surgical.select(deck, pages, self.root / 'never.pptx')
        self.assertFalse((self.root / 'never.pptx').exists())


if __name__ == '__main__':
    unittest.main()
