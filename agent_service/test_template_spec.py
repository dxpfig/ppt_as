"""Template extraction step 2: the rendered overview the design is read from."""
import json
from pathlib import Path
import re
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from agent_service import template_spec
from agent_service import template_kit
from agent_service.pptd_store import _serialize


def sample(path: Path) -> Path:
    prs = Presentation()
    for index in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Pt(60), Pt(60), Pt(800), Pt(60)).text_frame.text = f'第 {index + 1} 页'
    prs.save(path)
    return path


class OverviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_montage_labels_every_page_and_scales_to_one_width(self):
        thumbs = [Image.new('RGB', (960, 540), '#ffffff') for _ in range(5)]
        sheet = template_spec._montage(thumbs, 4)
        self.assertEqual(sheet.width, 4 * template_spec.THUMB_WIDTH + 5 * template_spec.MARGIN)
        rows = 2
        self.assertEqual(sheet.height,
                         rows * (round(540 * template_spec.THUMB_WIDTH / 960) + template_spec.CAPTION)
                         + (rows + 1) * template_spec.MARGIN)

    def test_rendered_overview_lists_pages_and_the_sheet(self):
        source = sample(self.root / 'deck.pptx')
        try:
            import asyncio
            result = asyncio.run(template_spec.overview(source, self.root / 'overview', columns=3))
        except template_spec.OverviewUnavailable as error:
            self.skipTest(f'渲染器不可用：{error}')
        self.assertEqual(result['pageCount'], 3)
        self.assertEqual(result['pages'], ['page-01.png', 'page-02.png', 'page-03.png'])
        sheet = self.root / 'overview' / result['overview']
        self.assertTrue(sheet.is_file())
        self.assertGreater(sheet.stat().st_size, 1000)

    def test_missing_source_is_reported_not_crashed(self):
        import asyncio
        with self.assertRaises(template_spec.OverviewUnavailable):
            asyncio.run(template_spec.overview(self.root / 'nope.pptx', self.root / 'out'))

    def test_model_attachments_put_the_sheet_first(self):
        folder = self.root / 'overview'
        folder.mkdir()
        (folder / 'overview.png').write_bytes(b'x')
        for index in range(1, 9):
            (folder / f'page-{index:02d}.png').write_bytes(b'x')
        attachments = template_spec.read_overview(folder)
        self.assertEqual(attachments[0]['role'], 'template-overview')
        self.assertEqual(len(attachments), 7)  # sheet + six key pages

    def test_tokens_read_the_design_the_pages_paint(self):
        def text(identity, value, size, color, y, x=40):
            return {'elementId': identity, 'elementType': 'text', 'bounds': [x, y, 600, 30],
                    'content': {'text': f'<p><span style="font-family:微软雅黑;font-size:{size}px;color:{color}">{value}</span></p>'}}

        pages = [
            {'elements': [text('cover', '申购汇报', 66, '#FFFFFF', 200)]},
            {'elements': [text('t2', '一、必要性', 30, '#E50011', 20),
                          text('b2', '正文内容', 13, '#000000', 120),
                          text('f2', '公司名', 12, '#333333', 500)]},
            {'elements': [text('t3', '二、技术路线', 30, '#E50011', 20),
                          text('b3', '正文内容', 13, '#000000', 120),
                          text('f3', '公司名', 12, '#333333', 500)]},
            {'elements': [text('t4', '三、预算', 30, '#E50011', 20),
                          text('b4', '正文内容', 13, '#000000', 120),
                          text('f4', '公司名', 12, '#333333', 500)]},
            {'elements': [text('closing', '请领导讲话', 100, '#FFFFFF', 200)]},
        ]
        result = template_spec.tokens(pages)
        self.assertEqual(result['textStyles']['page-title']['fontSize'], 30.0)
        self.assertEqual(result['textStyles']['page-title']['color'], '#E50011')
        self.assertEqual(result['textStyles']['body']['fontSize'], 13.0)
        self.assertEqual(result['textStyles']['footer']['fontSize'], 12.0)
        self.assertEqual(result['colors']['primary'], '#E50011')
        self.assertEqual(result['metrics']['leftMargin'], 40.0)

    def test_components_keep_only_repeated_parts(self):
        def shape(identity, name, width, height):
            return {'elementId': identity, 'elementType': 'shape', 'shapeName': name,
                    'bounds': [10, 10, width, height]}
        pages = [{'elements': [shape(f'c{index}', 'roundRect', 163, 300)]} for index in range(3)]
        pages.append({'elements': [shape('once', 'star5', 80, 80)]})
        found = template_spec.components(pages)
        self.assertEqual([item['name'] for item in found], ['roundRect'])
        self.assertEqual(found[0]['uses'], 3)
        self.assertEqual(found[0]['size'], [163, 300])
        self.assertNotIn('star5', [item['name'] for item in found])

    def test_spec_schema_and_persistence(self):
        pass
        valid = {'summary': '一份企业模板', 'pageTypes': [
            {'id': 'cover', 'purpose': '开场', 'templatePages': [1], 'skeleton': '满版背景 + 大标题',
             'slots': [{'name': 'title', 'purpose': '标题', 'characters': 20, 'repeatable': False}]}],
            'components': [{'id': 'logo', 'kind': 'image', 'description': '右上角标识', 'evidence': [2, 3]}],
            'rules': ['标题 30px #E50011'], 'prohibitions': ['不使用蓝色']}
        spec = template_spec.TemplateSpec.model_validate(valid, strict=True)
        self.assertEqual(spec.pageTypes[0].slots[0].characters, 20)
        with self.assertRaises(ValueError):
            template_spec.TemplateSpec.model_validate({'summary': 'x', 'pageTypes': []}, strict=True)
        with self.assertRaises(ValueError):
            template_spec.TemplateSpec.model_validate(
                {**valid, 'pageTypes': [{**valid['pageTypes'][0], 'unknown': 1}]}, strict=True)
        folder = self.root / 'spec'
        folder.mkdir()
        (folder / 'spec.json').write_text(json.dumps(valid, ensure_ascii=False), encoding='utf-8')
        self.assertEqual(template_spec.load_spec(folder)['summary'], '一份企业模板')
        self.assertIsNone(template_spec.load_spec(self.root / 'missing'))

    def test_repeating_unit_grows_shrinks_and_renumbers(self):
        def row(index, label):
            top = 160 + index * 44
            return [
                {'elementId': f'badge{index}', 'elementType': 'shape', 'shapeName': 'roundRect',
                 'bounds': [288, top, 62, 34], 'fill': {'type': 'solid', 'color': '#E60013'}},
                {'elementId': f'badge{index}-text', 'elementType': 'text', 'bounds': [288, top, 62, 34],
                 'content': {'text': f'<p><span style="font-size:26.0px">{index + 1:02d}</span></p>'}},
                {'elementId': f'label{index}', 'elementType': 'text', 'bounds': [354, top + 2, 104, 29],
                 'content': {'text': f'<p><span style="font-size:13.0px">条目{index + 1}</span></p>'}},
            ]
        page = {'elements': [item for index in range(4) for item in row(index, '')]
                + [{'elementId': 'title', 'elementType': 'text', 'bounds': [428, 88, 123, 56],
                    'content': {'text': '<p><span style="font-size:40.0px">目 录</span></p>'}}]}

        def numbers(value):
            return [re.sub(r'<[^>]+>', '', element['content']['text']) for element in value['elements']
                    if element.get('elementType') == 'text' and re.fullmatch(r'\d{2}', re.sub(r'<[^>]+>', '', element['content']['text']))]

        self.assertEqual(len(template_spec.repeat_units(page)['units']), 4)
        grown = template_spec.expand_repeats(page, 6)
        self.assertEqual(numbers(grown), ['01', '02', '03', '04', '05', '06'])
        self.assertEqual(len(grown['elements']), len(page['elements']) + 6)
        self.assertEqual([element for element in grown['elements'] if element['elementId'] == 'title'],
                         [element for element in page['elements'] if element['elementId'] == 'title'])
        shrunk = template_spec.expand_repeats(page, 2)
        self.assertEqual(numbers(shrunk), ['01', '02'])
        filled = template_spec.fill_slots(grown, {'repeat': ['一', '二', '三', '四', '五', '六']})
        self.assertIn('六', ' '.join(re.sub(r'<[^>]+>', '', element['content']['text'])
                                    for element in filled['elements'] if element.get('elementType') == 'text'))


def text_element(identity, bounds, value, *, size=30.0, color='#E50011'):
    return {'elementId': identity, 'elementType': 'text', 'bounds': list(bounds),
            'content': {'text': f'<p><span style="font-size:{size}px;color:{color}">{value}</span></p>',
                        'wrap': True, 'align': ['left', 'top']}}


def shape_element(identity, name, bounds, color='#E50011'):
    return {'elementId': identity, 'elementType': 'shape', 'shapeName': name, 'bounds': list(bounds),
            'fill': {'type': 'solid', 'color': color}}


def image_element(identity, bounds, src):
    return {'elementId': identity, 'elementType': 'image', 'bounds': list(bounds), 'src': src,
            'fit': {'mode': 'fill'}}


def cover_page():
    return {'size': [960, 540], 'elements': [
        image_element('s1', [0.0, 0.0, 958.0, 538.0], 'media/image-1.jpg'),
        text_element('s2', [140, 220, 680, 90], '信息化项目汇报', size=66.0, color='#FFFFFF'),
        text_element('s3', [140, 420, 300, 30], '汇报人：张三  2026-10-12', size=19.0, color='#FFFFFF')]}


def content_page(extra=None):
    """A body page with the chrome every rebuilt page has to keep."""
    return {'size': [960, 540], 'elements': [
        image_element('s2', [0.1, 0.6, 960.0, 539.9], 'media/image-5.jpg'),
        image_element('s3', [758.0, -22.4, 210.1, 129.7], 'media/image-4.png'),
        shape_element('s4', 'rect', [-0.3, 49.8, 31.8, 21.8]),
        text_element('s6', [64.7, 22.5, 620.0, 36.3], '方向一：示例标题'),
        text_element('s7', [64.7, 82.0, 850.0, 34.0], '导语：示例导语', size=19.0, color='#333333'),
        shape_element('s34', 'custom', [56.0, 80.0, 851.0, 1.0], color='#BABABA'),
        image_element('s41', [729.5, 517.7, 73.7, 13.4], 'media/image-7.png'),
        text_element('s42', [804.7, 515.9, 155.5, 14.6], '长沙湘计海盾科技有限公司', size=12.0, color='#333333'),
        *(extra or [])]}


def toc_page(rows=3):
    elements = []
    for index in range(rows):
        top = 160 + index * 44
        elements += [
            shape_element(f'b{index}', 'roundRect', [288, top, 62, 34]),
            text_element(f'b{index}-t', [288, top, 62, 34], f'{index + 1:02d}', size=26.0),
            text_element(f'l{index}', [354, top + 2, 104, 29], f'条目{index + 1}', size=13.0),
            shape_element(f'bar{index}', 'rect', [354, top + 2, 560, 29], color='#F3D6D9')]
    return {'size': [960, 540], 'elements': [
        *elements, text_element('t', [428, 88, 123, 56], '目 录', size=40.0)]}


def closing_page():
    return {'size': [960, 540], 'elements': [
        image_element('s1', [0.0, 0.0, 958.0, 538.0], 'media/image-2.jpg'),
        text_element('s2', [200, 240, 560, 80], '谢谢聆听', size=100.0, color='#FFFFFF')]}


class RebuildTests(unittest.TestCase):
    """Rebuild route: cover/contents/closing stay faithful, body pages compose."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        # 1 cover | 2-3 content | 4 contents | 5 content | 6 closing
        self.pages = [cover_page(), content_page(), content_page(), toc_page(),
                      content_page(), closing_page()]
        self.spec = {'summary': '示例模板', 'pageTypes': [
            {'id': 'cover', 'purpose': '封面', 'templatePages': [1], 'skeleton': '整版背景 + 大标题',
             'slots': [{'name': 'cover-title', 'purpose': '标题', 'characters': 20, 'repeatable': False},
                       {'name': 'cover-subtitle', 'purpose': '副标题', 'characters': 30, 'repeatable': False}]},
            {'id': 'content-sectioned', 'purpose': '正文', 'templatePages': [2, 3, 5], 'skeleton': '标题 + 分区',
             'slots': [{'name': 'page-title', 'purpose': '页标题', 'characters': 18, 'repeatable': False},
                       {'name': 'section-body', 'purpose': '正文', 'characters': 80, 'repeatable': True}]},
            {'id': 'toc', 'purpose': '目录', 'templatePages': [4], 'skeleton': '编号方块 + 白条',
             'slots': [{'name': 'toc-title', 'purpose': '目录标题', 'characters': 4, 'repeatable': False},
                       {'name': 'toc-entry-label', 'purpose': '条目', 'characters': 6, 'repeatable': True}]},
            {'id': 'closing', 'purpose': '结束页', 'templatePages': [6], 'skeleton': '整版背景 + 大字',
             'slots': [{'name': 'closing-title', 'purpose': '大字', 'characters': 10, 'repeatable': False}]}],
            'components': [], 'rules': [], 'prohibitions': []}
        self.summary = {'design': {
            'textStyles': {'page-title': {'fontSize': 30.0, 'fontFamily': 'Noto Sans SC', 'color': '#E50011'},
                           'lead': {'fontSize': 19.0, 'fontFamily': 'Noto Sans SC', 'color': '#333333'},
                           'body': {'fontSize': 13.0, 'fontFamily': 'Noto Sans SC', 'color': '#333333'}},
            'colors': {'primary': '#E50011'}}, 'components': []}

    def tearDown(self):
        self.tmp.cleanup()

    def kit(self):
        return template_kit.extract(self.pages, self.summary['design'], self.summary['components'])

    def test_only_cover_contents_and_closing_may_reuse_a_template_page(self):
        self.assertEqual(template_spec.faithful_page_types(self.spec), {'cover', 'toc', 'closing'})
        plan = {'title': 'T', 'pages': [
            {'pageType': 'content-sectioned', 'templatePage': 2, 'fills': {'page-title': '目标'}}]}
        issues = template_spec.validate_plan(plan, self.spec, self.summary)
        self.assertEqual([issue['code'] for issue in issues], ['BodyPageMustCompose'])

    def test_contents_rows_may_be_longer_than_the_example_wording(self):
        plan = {'title': 'T', 'pages': [
            {'pageType': 'toc', 'templatePage': 4, 'fills': {'toc-title': '目录'},
             'entries': ['项目建设背景与目标', '总体架构设计', '实施进度与里程碑', '投资预算与效益分析']}]}
        self.assertEqual(template_spec.validate_plan(plan, self.spec, self.summary), [])

    def test_a_longer_outline_continues_on_another_contents_page(self):
        # The template's list stops above the footer, so an outline with more
        # sections than that continues on a second copy of the same page and the
        # rows keep counting where the first page stopped.
        page = toc_page()
        self.assertEqual(template_spec.repeat_capacity(page), 7)
        entries = [f'第{i + 1}章' for i in range(9)]
        plan = {'title': 'T', 'pages': [
            {'pageType': 'toc', 'templatePage': 4, 'fills': {'toc-title': '目录'},
             'entries': entries}]}
        pages = [cover_page(), content_page(), content_page(), page, content_page(), closing_page()]
        issues = template_spec.validate_plan(plan, self.spec, self.summary, None,
                                             pages)
        self.assertEqual(issues, [])
        built = template_spec.build(pages, self.spec, plan, self.kit())
        self.assertEqual(len(built['pages']), 2)
        labels = [template_kit.plain(element.get('content', {}).get('text'))
                  for element in built['pages'][0]['data']['elements'] if element.get('elementType') == 'text']
        self.assertIn('第1章', labels)
        self.assertIn('第7章', labels)
        second = [template_kit.plain(element.get('content', {}).get('text'))
                  for element in built['pages'][1]['data']['elements'] if element.get('elementType') == 'text']
        self.assertIn('第8章', second)
        self.assertIn('第9章', second)
        self.assertIn('08', second)
        self.assertIn('09', second)

    def test_repeated_template_page_is_copied_not_aliased(self):
        plan = {'title': '信息化项目汇报', 'pages': [
            {'pageType': 'cover', 'templatePage': 1,
             'fills': {'cover-title': '信息化项目汇报', 'cover-subtitle': '汇报人：张三  2026-10-12'}},
            {'pageType': 'toc', 'templatePage': 4, 'fills': {'toc-title': '目录'},
             'entries': ['建设背景', '总体设计', '实施计划', '效益分析']},
            {'layout': 'cards', 'title': '建设背景', 'items': [{'label': '统一入口', 'bullets': ['打通系统']}]},
            {'pageType': 'cover', 'templatePage': 1,
             'fills': {'cover-title': '信息化项目汇报', 'cover-subtitle': '汇报人：李四  2026-11-01'}},
            {'pageType': 'closing', 'templatePage': 6, 'fills': {'closing-title': '谢谢聆听'}}]}
        built = template_spec.build(self.pages, self.spec, plan, self.kit())
        self.assertEqual(len(built['pages']), 5)
        for page in built['pages']:
            text = _serialize(page['data'])          # aliases make the store refuse the file
            self.assertNotIn('&id', text)
        first, second = built['pages'][0]['data'], built['pages'][3]['data']
        self.assertIsNot(first, second)
        self.assertFalse({id(element) for element in first['elements']}
                         & {id(element) for element in second['elements']})

    def test_composed_body_page_carries_the_template_chrome_once(self):
        kit = self.kit()
        self.assertEqual(kit['titleStyle'], 'page-title')
        plan = {'layout': 'cards', 'title': '建设背景与目标', 'lead': '围绕业务痛点明确四项目标',
                'items': [{'label': '统一入口', 'bullets': ['打通现有分散系统']},
                          {'label': '数据贯通', 'bullets': ['建立主数据标准']}]}
        page = template_kit.compose_body(kit, plan, page_id='p03')
        sources = [element.get('src') for element in page['elements'] if element.get('elementType') == 'image']
        self.assertEqual(sources.count('media/image-5.jpg'), 1)   # background, once
        self.assertEqual(sources.count('media/image-4.png'), 1)   # top-right mark, once
        self.assertEqual(sources.count('media/image-7.png'), 1)   # bottom-right mark, once
        decorations = [element for element in page['elements'] if element['elementId'].startswith('p03-deco')]
        self.assertEqual([element['elementId'] for element in decorations], ['p03-deco0', 'p03-deco1'])
        self.assertEqual([element['elementType'] for element in decorations], ['shape', 'shape'])
        title = next(element for element in page['elements'] if element['elementId'] == 'p03-title')
        self.assertIn('font-size:30.0px', title['content']['text'])
        self.assertIn('color:#E50011', title['content']['text'])
        self.assertIn('建设背景与目标', title['content']['text'])
        lead = next(element for element in page['elements'] if element['elementId'] == 'p03-lead')
        self.assertIn('font-size:19.0px', lead['content']['text'])

    def test_composed_item_without_words_is_rejected(self):
        plan = {'title': 'T', 'pages': [
            {'layout': 'cards', 'title': '建设背景',
             'items': [{'title': '统一入口', 'body': '打通现有分散系统'}]}]}
        issues = template_spec.validate_plan(plan, self.spec, self.summary)
        self.assertEqual([issue['code'] for issue in issues], ['EmptyItem'])
        self.assertIn('label', issues[0]['message'])

    def test_rewriting_a_slot_keeps_its_paragraphs_and_fits_its_box(self):
        # The company cover holds its title up with a trailing empty paragraph, so
        # rewriting must keep every tag and give the box the words' own metrics.
        content = {'text': '<p style="text-align:center;line-height:1.2">'
                           '<span style="font-family:MiSans;font-size:66.0px">关于</span></p>'
                           '<p style="text-align:center;line-height:1.2"></p>',
                   'wrap': True, 'align': ['left', 'middle']}
        rewritten = template_kit.rewrite(content, '公司级AI能力平台建设')
        self.assertEqual(rewritten['text'].count('<p'), 2)
        self.assertIn('&#160;', rewritten['text'], 'the blank line the template relied on must survive')
        self.assertIn('text-align:center', rewritten['text'])
        self.assertIn('公司级AI能力平台建设', rewritten['text'])
        self.assertEqual(rewritten['fontSize'], 66.0)
        self.assertEqual(rewritten['lineHeight'], template_kit.LINE_SPACING)
        element = {'elementId': 's1', 'elementType': 'text', 'bounds': [0.0, 0.0, 960.0, 36.0],
                   'content': rewritten}
        template_kit.fit(element, {'size': [960, 540]})
        self.assertGreaterEqual(element['bounds'][3], template_kit.LINE_SPACING * 66.0)

    def test_a_narrow_slot_widens_for_a_one_line_phrase(self):
        content = {'text': '<p style="text-align:center">'
                           '<span style="font-family:微软雅黑;font-size:18.0px">谢谢聆听</span></p>'}
        element = {'elementId': 's2', 'elementType': 'text', 'bounds': [400.0, 300.0, 174.7, 46.0],
                   'content': content}
        element['content'] = template_kit.rewrite(
            {**content, 'text': content['text'].replace('谢谢聆听', '公司级AI能力平台建设专题汇报')},
            '公司级AI能力平台建设专题汇报')
        template_kit.fit(element, {'size': [960, 540]})
        self.assertGreater(element['bounds'][2], 174.7)
        self.assertLess(element['bounds'][0], 400.0)
        self.assertGreater(element['bounds'][3], 18.0)

    def test_body_content_that_cannot_fit_its_box_is_reported(self):
        kit = self.kit()
        long_bullet = '打通现有分散系统并完成主数据标准建设与集成测试全部工作'
        plan = {'title': 'T', 'pages': [
            {'layout': 'steps', 'title': '实施计划',
             'items': [{'label': f'第{i}阶段目标', 'bullets': [long_bullet] * 3} for i in range(6)]}]}
        issues = template_spec.validate_plan(plan, self.spec, self.summary, kit)
        self.assertTrue(issues)
        self.assertTrue(all(issue['code'] == 'SlotOverflow' for issue in issues))

    def test_the_same_plan_fits_once_the_words_match_the_boxes(self):
        kit = self.kit()
        plan = {'title': 'T', 'pages': [
            {'layout': 'steps', 'title': '实施计划',
             'items': [{'label': f'第{i + 1}阶段', 'bullets': ['完成选型与部署']} for i in range(3)]}]}
        self.assertEqual(template_spec.validate_plan(plan, self.spec, self.summary, kit), [])

    def test_a_card_grid_fills_its_rows_and_centres_the_last_one(self):
        self.assertEqual(template_kit._columns(4), (1, 4))
        self.assertEqual(template_kit._columns(5), (2, 3))
        self.assertEqual(template_kit._columns(10), (3, 4))
        plan = {'layout': 'cards', 'title': '总体架构',
                'items': [{'label': f'第{i + 1}层', 'bullets': ['职责说明']} for i in range(5)]}
        page = template_kit.compose_body(self.kit(), plan, page_id='p01')
        cards = [element for element in page['elements']
                 if element['elementId'].startswith('p01-card') and element['elementId'].endswith('-bg')]
        self.assertEqual(len(cards), 5)
        first, last = cards[0]['bounds'], cards[-1]['bounds']
        # The two cards of the second row straddle the centre of the first row.
        centre = first[0] + (first[2] * 3 + 40) / 2
        second_row = [card['bounds'] for card in cards[3:]]
        self.assertAlmostEqual(second_row[0][0] + second_row[-1][0] + second_row[-1][2], centre * 2, delta=2)


if __name__ == '__main__':
    unittest.main()
