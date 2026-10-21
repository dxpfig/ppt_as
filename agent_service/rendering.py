"""Preserve corporate OOXML artwork and render actual PPTX through LibreOffice."""
import asyncio
from copy import deepcopy
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import subprocess
import unicodedata
import zipfile

import pymupdf
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from PIL import Image
from pptx.util import Inches, Pt

from .design import compose_slide, FONT_DIR
from .resource_manager import RENDER_LIMIT

MAX_UPLOAD = 30 * 1024 * 1024
DOCKER_IMAGE = os.getenv("PPT_RENDERER_IMAGE", "wechatopenai/weknora-docreader:latest")
RENDER_LOCK = RENDER_LIMIT


def renderer_available():
    if shutil.which(os.getenv("LIBREOFFICE_BIN", "libreoffice")):
        return True
    if not shutil.which("docker"):
        return False
    try:
        return subprocess.run(["docker", "image", "inspect", DOCKER_IMAGE], capture_output=True, timeout=5).returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


def validate_package(data: bytes):
    if len(data) > MAX_UPLOAD:
        raise ValueError("模板不能超过 30 MB。")
    try:
        archive = zipfile.ZipFile(BytesIO(data))
        entries = archive.infolist()
        if len(entries) > 3000 or sum(e.file_size for e in entries) > 200 * 1024 * 1024:
            raise ValueError("模板解压后过大或包含过多文件。")
        for entry in entries:
            path = entry.filename.lower()
            if ".." in Path(path).parts or path.startswith("/") or (entry.compress_size and entry.file_size / entry.compress_size > 1000):
                raise ValueError("模板压缩包结构不安全。")
            if "vbaproject" in path or "/activex/" in path or path.endswith(".exe") or ("/embeddings/" in path and not path.endswith(".xlsx")):
                raise ValueError("暂不支持包含宏、OLE 或嵌入程序的模板，请另存为普通 PPTX 后上传。")
            if "/embeddings/" in path and path.endswith(".xlsx"):
                with zipfile.ZipFile(BytesIO(archive.read(entry))) as workbook:
                    parts = workbook.infolist()
                    if len(parts) > 1000 or sum(p.file_size for p in parts) > 40 * 1024 * 1024:
                        raise ValueError("模板内嵌图表工作簿过大。")
                    for part in parts:
                        part_name = part.filename.lower()
                        if "vbaproject" in part_name or "/embeddings/" in part_name or "/activex/" in part_name or part_name.endswith((".bin", ".exe")):
                            raise ValueError("模板内嵌工作簿包含不支持的活动内容。")
                        if part_name.endswith((".xml", ".rels")):
                            content = workbook.read(part)
                            if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                                raise ValueError("模板内嵌工作簿包含不安全 XML。")
                            tree = etree.fromstring(content, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                            if part_name.endswith(".rels") and any(r.get("TargetMode") == "External" for r in tree):
                                raise ValueError("模板内嵌工作簿包含外部链接，请移除后上传。")
            if path.endswith((".xml", ".rels")):
                contents = archive.read(entry)
                if b"<!DOCTYPE" in contents.upper() or b"<!ENTITY" in contents.upper():
                    raise ValueError("模板包含不支持的 XML 声明。")
                try:
                    tree = etree.fromstring(contents, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                except etree.XMLSyntaxError as exc:
                    raise ValueError("模板 XML 内容损坏，请另存为新的 PPTX 后重试。") from exc
                if path.endswith(".rels"):
                    for relationship in tree:
                        if relationship.get("TargetMode") != "External" or relationship.get("Type", "").endswith("/hyperlink"):
                            continue
                        # PowerPoint leaves a dangling chart externalData/oleObject
                        # relationship with an empty target after the workbook link
                        # is stripped. No external resource exists to fetch, and the
                        # chart values stay cached in the chart part.
                        if not (relationship.get("Target") or "").strip():
                            continue
                        raise ValueError("模板含外部链接资源，请先在 PowerPoint 中嵌入图片/媒体后再上传。")
        if "ppt/presentation.xml" not in archive.namelist():
            raise ValueError("文件不是有效的 PowerPoint 模板。")
        # python-pptx expects the presentation MIME type; .potx is otherwise identical.
        out = BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as normalized:
            for entry in entries:
                content = archive.read(entry)
                if entry.filename == "[Content_Types].xml":
                    content = content.replace(b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml", b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")
                normalized.writestr(entry.filename, content)
        return out.getvalue()
    except zipfile.BadZipFile as exc:
        raise ValueError("模板文件损坏，请上传有效的 .pptx 或 .potx。") from exc


def template_info(path: Path):
    prs = Presentation(path)
    if len(prs.slides) > 60:
        raise ValueError("模板最多包含 60 页，请仅保留需要的示例页。")
    warnings = ["保留模板母版、主题、背景、图片及文本框样式；按示例页顺序套用，超出示例数量后复用最后一页。", "图表、表格及装饰文字会保留原内容；生成的 PPTX、静态预览及 PDF 均不保留原模板动画和切换效果。"]
    return {"slideCount": len(prs.slides), "layouts": [layout.name for layout in prs.slide_layouts], "warnings": warnings}


def replace_text(shape, lines):
    frame = shape.text_frame
    prototypes = [deepcopy(p._p) for p in frame.paragraphs]
    for paragraph in list(frame._txBody)[2:]:
        frame._txBody.remove(paragraph)
    for index, line in enumerate(lines or [""]):
        p = deepcopy(prototypes[min(index, len(prototypes) - 1)])
        runs = p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
        run_style = deepcopy(runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")) if runs else None
        for node in list(p):
            if etree.QName(node).localname not in ("pPr", "endParaRPr"):
                p.remove(node)
        run = etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
        if run_style is not None:
            run.append(run_style)
        etree.SubElement(run, "{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = line
        p.insert(1 if len(p) and etree.QName(p[0]).localname == "pPr" else 0, run)
        frame._txBody.append(p)


def clone_slide(prs, source):
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    mapping = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(("/notesSlide", "/slideLayout")):
            continue
        mapping[rel.rId] = slide.part.relate_to(rel.target_ref if rel.is_external else rel.target_part, rel.reltype, rel.is_external)
    for shape in source.shapes:
        copied = deepcopy(shape._element)
        for node in copied.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in mapping:
                    node.set(attr, mapping[value])
        slide.shapes._spTree.insert_element_before(copied, "p:extLst")
    background = source._element.cSld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if background is not None:
        copied = deepcopy(background)
        for node in copied.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in mapping:
                    node.set(attr, mapping[value])
        slide._element.cSld.insert(0, copied)
    return slide


def populate_template_slide(slide, content, prs, warnings):
    texts = [s for s in slide.shapes if s.has_text_frame and (not s.is_placeholder or s.placeholder_format.type not in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER))]
    title = slide.shapes.title
    if title is None and texts:
        title = min(texts, key=lambda s: s.top)
    if title is not None:
        replace_text(title, [content.title])
    else:
        title = slide.shapes.add_textbox(Inches(.7), Inches(.5), prs.slide_width - Inches(1.4), Inches(.8))
        title.text = content.title
        warnings.append("部分模板页缺少标题框，已补充标题框。")
    others = [s for s in texts if s._element is not title._element]
    subtitle = next((s for s in others if s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE), None)
    if subtitle is not None:
        replace_text(subtitle, [content.subtitle])
        others.remove(subtitle)
    body = max(others, key=lambda s: s.width * s.height, default=None)
    lines = ([content.subtitle] if content.subtitle and subtitle is None else []) + content.points
    if body is None:
        body = slide.shapes.add_textbox(Inches(.8), Inches(1.6), prs.slide_width - Inches(1.6), prs.slide_height - Inches(2))
        body.text_frame.paragraphs[0].font.size = Pt(22)
        warnings.append("部分模板页缺少正文框，已按模板主题补充正文框；请检查排版。")
    replace_text(body, lines)
    if len(others) > 1:
        warnings.append("模板包含多个文本框，仅替换标题和主要正文，其他装饰/示例文字保留，请检查。")


def place_picture(slide, path, x, y, width, height, content):
    # Contain rather than crop: knowledge-base diagrams must remain readable in full.
    with Image.open(path) as image:
        ratio = min(width / image.width, height / image.height)
        w, h = image.width * ratio, image.height * ratio
    picture = slide.shapes.add_picture(str(path), Inches(x + (width-w)/2), Inches(y + (height-h)/2), width=Inches(w), height=Inches(h))
    picture._element.nvPicPr.cNvPr.set("descr", content.image.title if content.image else content.title)
    if content.image and content.image.sourceUrl.startswith(("https://", "http://")):
        picture.click_action.hyperlink.address = content.image.sourceUrl


def picture_notes(content, used):
    notes = content.notes
    if content.source:
        notes += "\n\n内容来源：" + content.source
    if used and content.image:
        image = content.image
        notes += f"\n\n配图：{image.title}\n来源：{image.provider} / {image.sourceUrl or '本地知识库'}\n作者：{image.author or '来源未标注'}\n许可：{image.license}"
    return notes


def create_deck(project, path: Path, template_path: Path | None, images: dict | None = None):
    images = images or {}
    warnings = []
    if template_path:
        prs = Presentation(template_path)
        warnings.extend(template_info(template_path)["warnings"])
        sources = list(prs.slides)
        old_ids = list(prs.slides._sldIdLst)
        for index, content in enumerate(project.slides):
            slide = clone_slide(prs, sources[min(index, len(sources)-1)]) if sources else prs.slides.add_slide(prs.slide_layouts[min(1, len(prs.slide_layouts)-1)])
            populate_template_slide(slide, content, prs, warnings)
            used = False
            if content.id in images:
                placeholder = next((shape for shape in slide.placeholders if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE and hasattr(shape, "insert_picture")), None)
                if placeholder is not None:
                    # Remove the empty placeholder only; preserve all company artwork and text positions.
                    x, y, width, height = (v / 914400 for v in (placeholder.left, placeholder.top, placeholder.width, placeholder.height))
                    placeholder._element.getparent().remove(placeholder._element)
                    place_picture(slide, images[content.id], x, y, width, height, content)
                    used = True
                else:
                    warnings.append(f"第 {index+1} 页模板没有图片占位符，配图已保留在编辑区，未覆盖公司模板内容。")
            slide.notes_slide.notes_text_frame.text = picture_notes(content, used)
        for slide_id in old_ids:
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for index, content in enumerate(project.slides):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            page_warnings, used = compose_slide(slide, content, project, index, images.get(content.id), place_picture)
            warnings.extend(page_warnings)
            slide.notes_slide.notes_text_frame.text = picture_notes(content, used)
    for content in project.slides:
        if sum(map(len, content.points)) > 600 or len(content.title) > 45:
            warnings.append("部分页面文字较多，可能超出模板文本框，请在预览中检查并精简内容。")
    prs.save(path)
    return list(dict.fromkeys(warnings))


async def render_pdf(path: Path):
    async with RENDER_LOCK:
        binary = shutil.which(os.getenv("LIBREOFFICE_BIN", "libreoffice"))
        container = f"ppt-studio-{path.parent.name}"
        if binary:
            args = [binary, f"-env:UserInstallation=file://{path.parent / 'lo-profile'}", "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)]
        elif renderer_available():
            args = ["docker", "run", "--rm", "--name", container, "--network", "none", "--cap-drop", "ALL", "--security-opt", "no-new-privileges", "--memory", "1g", "--cpus", "2", "--user", f"{os.getuid()}:{os.getgid()}", "-e", "HOME=/tmp", "-v", f"{path.parent.resolve()}:/work:rw", "--entrypoint", "libreoffice", DOCKER_IMAGE, "-env:UserInstallation=file:///tmp/lo-profile", "--headless", "--convert-to", "pdf", "--outdir", "/work", f"/work/{path.name}"]
            if FONT_DIR.is_dir():
                args[2:2] = ["-v", f"{FONT_DIR}:/usr/share/fonts/opentype/pptstudio:ro"]
        else:
            raise ValueError("未找到 LibreOffice 渲染器，无法生成真实预览/PDF。请安装 LibreOffice 或配置 PPT_RENDERER_IMAGE。")
        proc = await asyncio.create_subprocess_exec(*args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
        try:
            await asyncio.wait_for(proc.communicate(), timeout=150)
            if proc.returncode != 0 or not path.with_suffix(".pdf").exists():
                raise ValueError("LibreOffice 无法渲染此文件，请检查模板是否损坏或含有不支持的元素。")
        finally:
            if proc.returncode is None:
                proc.kill()
                await proc.wait()
                if not binary:
                    cleanup = await asyncio.create_subprocess_exec("docker", "rm", "-f", container, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
                    await cleanup.wait()
        return await asyncio.to_thread(rasterize, path.with_suffix(".pdf"))


def rasterize(pdf: Path):
    names = []
    with pymupdf.open(pdf) as document:
        for index, page in enumerate(document):
            name = f"slide-{index+1}.png"
            page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(pdf.parent / name)
            names.append(name)
    return names


def audit_rendered_pdf(project, pdf: Path):
    """Check actual output for page loss, off-page text and missing content.

    Text-layer checks are a warning, not a proof of visual correctness. Template
    reading order and font substitution can differ, so the PNG remains authoritative.
    """
    def normalized(text):
        return "".join(c for c in unicodedata.normalize("NFKC", text) if c.isalnum())
    warnings = []
    with pymupdf.open(pdf) as document:
        if len(document) != len(project.slides):
            warnings.append("实际预览页数与内容页数不一致，请检查导出文件。")
        for index, (page, content) in enumerate(zip(document, project.slides), 1):
            visible = normalized(page.get_text())
            expected = [content.title, content.subtitle, *content.points]
            if any(normalized(text) and normalized(text) not in visible for text in expected):
                warnings.append(f"第 {index} 页部分文字未在 PDF 文本层完整检出，请检查是否溢出或被截断。")
            blocks = page.get_text("dict").get("blocks", [])
            for block in blocks:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        x0, y0, x1, y1 = span["bbox"]
                        if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                            warnings.append(f"第 {index} 页文字接近或超出页面边界，请检查预览。")
    return list(dict.fromkeys(warnings))
