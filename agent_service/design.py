"""Editable slide compositions with measured typography and a shared design system."""
from functools import lru_cache
from pathlib import Path
import re
import unicodedata

from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

FONT_DIR = Path("/usr/share/fonts/opentype/noto")
FONT_FAMILY = "Noto Sans CJK SC" if FONT_DIR.is_dir() else "WenQuanYi Zen Hei"
PALETTES = {
    "forest": {"paper": "F5F3EC", "ink": "183E34", "muted": "62776F", "rule": "CCD3C9", "accent": "9B783E"},
    "ink": {"paper": "F1F2F4", "ink": "25344B", "muted": "617080", "rule": "CBD1D7", "accent": "9D643D"},
    "clay": {"paper": "FAF3E9", "ink": "683A2F", "muted": "8B7369", "rule": "DBCABF", "accent": "A75B35"},
}


@lru_cache(maxsize=96)
def measurement_font(size, bold=False):
    path = FONT_DIR / ("NotoSansCJK-Bold.ttc" if bold else "NotoSansCJK-Regular.ttc")
    try:
        return ImageFont.truetype(str(path), size * 4, index=2)
    except OSError:
        return None


def text_width(text, size, bold=False):
    font = measurement_font(size, bold)
    if font:
        return font.getlength(text) / 4
    return sum(size if unicodedata.east_asian_width(c) in "WF" else size * .58 for c in text)


def wrap_text(text, width, size, bold=False):
    """Wrap Latin words together and CJK characters by glyph, preserving all text."""
    lines = []
    for paragraph in text.split("\n"):
        current = ""
        for token in re.findall(r"[A-Za-z0-9_./:@%+\-]+|[^A-Za-z0-9_./:@%+\-]", paragraph):
            if current and text_width(current + token, size, bold) > width:
                carry = current[-1] if token[0] in "，。！？；：、）】》…’”,.!?;:)]}%" or current[-1] in "（【《‘“([{" else ""
                lines.append((current[:-1] if carry else current).rstrip())
                current = carry
            # Long URLs/identifiers can wrap even when they have no word boundary.
            for char in token:
                if current and text_width(current + char, size, bold) > width:
                    lines.append(current.rstrip())
                    current = ""
                current += char
        lines.append(current.rstrip())
    return "\n".join(lines)


def fit_text(text, width, height, preferred, minimum, bold=False):
    for size in range(preferred, minimum - 1, -1):
        wrapped = wrap_text(text, width * 72 - 4, size, bold)
        if len(wrapped.split("\n")) * size * 1.32 <= height * 72 - 3:
            return wrapped, size, True
    return wrapped, size, False


def effective_layout(content, index, total, has_image):
    layout = content.layout
    if layout == "auto":
        layout = "cover" if index == 0 and total > 1 else "image" if has_image else "editorial"
    if layout == "image" and not has_image:
        return "editorial"
    if layout == "image":
        required = sum(len(wrap_text(p, (5.5-.64)*72-4, 16).split("\n")) * 16 * 1.32 / 72 for p in content.points)
        if required + .14 * max(0, len(content.points)-1) + .08 * len(content.points) > 2.97:
            return "editorial"
    if layout == "comparison" and len(content.points) != 2:
        return "editorial"
    if layout == "process" and not 2 <= len(content.points) <= 5:
        return "editorial"
    if layout in ("cover", "closing") and len(content.points) > 3:
        return "image" if has_image else "editorial"
    return layout


class Canvas:
    def __init__(self, slide, index, palette, dark=False):
        self.slide, self.index, self.palette = slide, index, palette
        self.dark = dark
        self.foreground = palette["paper"] if dark else palette["ink"]
        self.muted = "C9D6CD" if dark else palette["muted"]
        self.rule_color = "62776F" if dark else palette["rule"]
        self.warnings = []
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(palette["ink"] if dark else palette["paper"])

    def text(self, x, y, w, h, text, size=22, minimum=16, bold=False, color=None, name="text"):
        if not text:
            return None
        wrapped, actual, fits = fit_text(text, w, h, size, minimum, bold)
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = f"ppt-{name}"
        frame = shape.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        # Soft line breaks retain one paragraph and a single style.
        frame.paragraphs[0].text = wrapped.replace("\n", "\v")
        for p in frame.paragraphs:
            p.font.name = FONT_FAMILY
            p.font.size, p.font.bold = Pt(actual), bold
            p.font.color.rgb = RGBColor.from_string(color or self.foreground)
            p.line_spacing = Pt(actual * 1.32)
            p.space_before = p.space_after = Pt(0)
            p._p.get_or_add_pPr().set("marL", "0")
            # Explicit East Asian typeface avoids theme font substitution in Office/LO.
            rpr = p._p.get_or_add_pPr().get_or_add_defRPr()
            from lxml import etree
            ea = etree.SubElement(rpr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
            ea.set("typeface", FONT_FAMILY)
        if not fits:
            self.warnings.append(f"第 {self.index+1} 页“{name}”文字超出推荐容量，请精简或拆页；内容已完整保留。")
        return shape

    def line(self, x1, y1, x2, y2, width=.7, color=None):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = RGBColor.from_string(color or self.rule_color)
        shape.line.width = Pt(width)
        shape.shadow.inherit = False
        # LibreOffice also reads the theme effectRef even with an empty effectLst.
        for effect in shape._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        return shape

    def rows(self, points, x, y, w, h, size=23):
        if not points:
            return
        available = h - .14 * max(0, len(points)-1) - .08 * len(points)
        for target in range(size, 15, -1):
            needed = [len(wrap_text(point, (w-.64)*72-4, target).split("\n")) * target * 1.32 / 72 + .055 for point in points]
            if sum(needed) <= available:
                break
        if sum(needed) <= available:
            heights = [value + (available-sum(needed))/len(points) for value in needed]
        else:
            heights = [max(.12, available * value/sum(needed)) for value in needed]
        cursor = y
        for i, (point, text_h) in enumerate(zip(points, heights)):
            row_h = text_h + .08
            self.text(x, cursor+.03, .42, .32, f"{i+1:02d}", 11, 10, color=self.muted, name="number")
            self.text(x+.64, cursor, w-.64, text_h, point, target, 16, name=f"point-{i+1}")
            cursor += row_h
            if i < len(points)-1:
                self.line(x+.64, cursor, x+w, cursor)
                cursor += .14


def compose_slide(slide, content, project, index, image_path, add_picture):
    total = len(project.slides)
    has_image = image_path is not None
    layout = effective_layout(content, index, total, has_image)
    c = Canvas(slide, index, PALETTES[project.theme], dark=layout in ("cover", "closing"))
    if content.layout not in ("auto", layout):
        c.warnings.append(f"第 {index+1} 页已根据图片和条目数量调整为可容纳完整内容的版式。")
    tag = content.section or {"cover": "开场", "editorial": "重点阐述", "image": "场景与证据", "comparison": "选项比较", "process": "实施路径", "closing": "下一步"}[layout]
    c.text(.65, .47, 9.6, .3, tag, 11, 10, color=c.muted, name="section")
    c.text(11.9, .47, .8, .3, f"{index+1:02d} / {total:02d}", 10, 10, color=c.muted, name="page-number")
    c.line(.65, .97, 12.68, .97)

    def picture(x, y, w, h):
        if has_image:
            add_picture(slide, image_path, x, y, w, h, content)

    if layout == "cover":
        width = 6.2 if has_image else 11.0
        c.text(.7, 1.48, width, 2.0, content.title, 44, 28, True, name="title")
        c.text(.73, 3.74, width, .85, content.subtitle, 21, 16, color=c.muted, name="subtitle")
        c.line(.73, 4.85, 1.65, 4.85, 3, c.palette["accent"])
        c.rows(content.points, .73, 5.08, width, 1.64, 18)
        picture(7.45, 1.45, 5.18, 5.13)
    elif layout == "closing":
        c.text(.7, 1.57, 5.2, 2.3, content.title, 38, 26, True, name="title")
        c.text(.73, 4.14, 4.85, 1.25, content.subtitle, 20, 16, color=c.muted, name="subtitle")
        if has_image:
            picture(7.0, 1.47, 5.63, 2.45)
            c.rows(content.points, 6.87, 4.23, 5.76, 2.39, 20)
        else:
            c.rows(content.points, 6.65, 1.7, 5.95, 4.65, 25)
    else:
        c.text(.7, 1.25, 11.9, 1.10, content.title, 32, 24, True, name="title")
        if layout == "image":
            c.text(.74, 2.62, 5.5, .94, content.subtitle, 19, 16, color=c.muted, name="subtitle")
            c.rows(content.points, .74, 3.65, 5.5, 2.97, 21)
            c.line(6.47, 2.64, 6.47, 6.57)
            picture(6.74, 2.60, 5.88, 3.99)
        elif layout == "comparison":
            c.text(.74, 2.52, 11.7, .65, content.subtitle, 18, 16, color=c.muted, name="subtitle")
            c.line(6.62, 3.6, 6.62, 6.56)
            for i, point in enumerate(content.points):
                x, width = (.76, 5.23) if i == 0 else (7.12, 5.46)
                heading, separator, body = re.sub(r"^\s*[•\-]\s*", "", point).partition("：")
                if not separator:
                    heading, separator, body = point.partition(":")
                c.text(x, 3.55, width, .70, heading if separator else f"选项 {i+1:02d}", 25, 18, True, name=f"option-{i+1}")
                c.line(x, 4.55, x+.65, 4.55, 2, c.palette["accent"])
                c.text(x, 4.84, width, 1.76, body if separator else point, 22, 16, name=f"comparison-{i+1}")
            if has_image:
                # Explicit user-selected imagery remains visible without covering comparisons.
                c.warnings.append(f"第 {index+1} 页对比版式以比较内容为主，所选图片保留在编辑区；切换图文版式可插入。")
                has_image = False
        elif layout == "process":
            c.text(.74, 2.55, 3.13, 1.75, content.subtitle, 22, 16, color=c.muted, name="subtitle")
            c.line(4.42, 2.72, 4.42, 6.35, 1, c.palette["accent"])
            step_h = 3.99 / len(content.points)
            uniform_size = min(fit_text(point, 7.06, step_h-.15, 23, 16)[1] for point in content.points)
            for i, point in enumerate(content.points):
                y = 2.62 + i * step_h
                c.line(4.29, y+.17, 4.55, y+.17, 3, c.palette["accent"])
                c.text(4.87, y, .48, .4, f"{i+1:02d}", 13, 11, color=c.muted, name="step-number")
                c.text(5.51, y, 7.06, step_h-.15, point, uniform_size, 16, name=f"step-{i+1}")
            if has_image:
                picture(.74, 4.47, 3.05, 2.12)
        else:
            c.text(.74, 2.62, 3.16, 1.72 if has_image else 3.2, content.subtitle, 23, 17, color=c.muted, name="subtitle")
            if has_image:
                picture(.74, 4.55, 3.15, 2.04)
            c.line(4.3, 2.63, 4.3, 6.6)
            c.rows(content.points, 4.7, 2.64, 7.91, 3.99, 25)
    c.line(.65, 6.97, 12.68, 6.97)
    footer = content.source or project.title
    footer_box = c.text(.68, 7.09, 9.10, .24, footer, 9, 8, color=c.muted, name="source")
    if footer_box and content.source:
        match = re.search(r"https?://[^\s，；）]+", content.source)
        if match:
            footer_box.text_frame.paragraphs[0].runs[0].hyperlink.address = match[0]
    if has_image and content.image:
        c.text(9.95, 7.09, 2.72, .24, f"配图 · {content.image.provider}", 9, 8, color=c.muted, name="image-credit")
    return c.warnings, has_image
