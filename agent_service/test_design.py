"""Prompt-library and native-layout regression checks, isolated from user records."""
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .design import effective_layout, fit_text, wrap_text
from .models import Project, Slide
from .prompt_library import PromptLibrary, PROMPT_ROOT
from .rendering import create_deck


def example_project():
    return Project(id="design-check", title="让城市屋顶成为清洁能源入口", theme="forest", imageSource="none", slides=[
        Slide(id="cover", layout="cover", section="城市能源 / 方案讨论", title="让闲置屋顶\n产生长期价值", subtitle="从建筑条件出发，规划分布式光伏的实施路径", points=["先核查屋顶条件，再讨论投资与收益"], imageIntent="none"),
        Slide(id="editorial", layout="editorial", section="判断依据", title="先确认适用条件，再决定是否建设", subtitle="价值来自真实的建筑与用电条件，不能由装机规模单独判断。", points=["结构安全：先确认承载能力、屋面状态与剩余使用年限", "用电匹配：比较日间负荷与发电时段，评估自用空间", "责任明确：在方案阶段约定建设、运维和产权边界"], imageIntent="none"),
        Slide(id="image", layout="image", section="场景观察", title="屋顶空间与用电负荷需要一起评估", subtitle="实景用于说明条件，不能代替现场勘察。", points=["查看屋顶朝向、遮挡与可用面积", "核对建筑实际用电规律", "预留检修与消防通道"], imageIntent="photo"),
        Slide(id="comparison", layout="comparison", section="模式比较", title="合作模式应匹配资金与管理能力", subtitle="以下为常见模式概述；具体权责以合同与项目条件为准。", points=["自行投资：自主安排建设与运营，同时承担资金投入、设备维护和收益波动。", "合作运营：由合作方承担约定的建设运维职责，需要事先明确收益分配、退出与交接条件。"], imageIntent="none"),
        Slide(id="process", layout="process", section="推进路径", title="按验证节点推进，降低返工风险", subtitle="每一步完成后，确认条件再进入下一阶段。", points=["现场核查：检查结构、遮挡、用电与并网条件", "方案比选：比较技术路径、费用与责任边界", "试点建设：先验证施工和运行安排", "运行复盘：记录实际表现，再决定后续推广"], imageIntent="none"),
        Slide(id="closing", layout="closing", section="下一步", title="从一份可核查的\n屋顶清单开始", subtitle="把讨论转为有依据的项目准备。", points=["整理建筑资料与用电记录", "组织现场核查，列明待确认条件", "形成试点方案，再审议实施安排"], imageIntent="none"),
    ])


class DesignTests(unittest.TestCase):
    def test_prompt_snapshot_and_next_job_reload(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "prompts"
            shutil.copytree(PROMPT_ROOT, root)
            before = PromptLibrary(root)
            with (root / "writer.md").open("a") as f:
                f.write("\n新的写作约束。\n")
            after = PromptLibrary(root)
            self.assertNotEqual(before.version, after.version)
            self.assertNotIn("新的写作约束", before.system("writer"))
            self.assertIn("新的写作约束", after.system("writer"))
            (root / "editor.md").write_text("")
            with self.assertRaisesRegex(ValueError, "为空"):
                PromptLibrary(root)

    def test_layout_fallback_keeps_all_content(self):
        s = Slide(id="old", title="旧工作记录", points=["第一项", "第二项", "第三项"])
        self.assertEqual(effective_layout(s, 0, 1, False), "editorial")
        s.layout = "comparison"
        self.assertEqual(effective_layout(s, 0, 1, False), "editorial")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fallback.pptx"
            warnings = create_deck(Project(id="fallback", slides=[s]), path, None)
            text = "".join(shape.text for shape in Presentation(path).slides[0].shapes if shape.has_text_frame)
            self.assertTrue(all(point in text for point in s.points))
            self.assertTrue(any("调整" in warning for warning in warnings))

    def test_native_layouts_fit_and_keep_editable_content(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            photo = root / "photo.png"
            Image.new("RGB", (900, 600), "#c0ceba").save(photo)
            p = example_project()
            path = root / "all-layouts.pptx"
            warnings = create_deck(p, path, None, {"image": photo})
            self.assertEqual(warnings, [])
            prs = Presentation(path)
            self.assertEqual(len(prs.slides), 6)
            for page, content in zip(prs.slides, p.slides):
                text = "".join(s.text for s in page.shapes if s.has_text_frame).replace("\v", "").replace("\n", "")
                self.assertIn(content.title.replace("\n", ""), text)
                for point in content.points:
                    # Comparison renders the option name and body separately.
                    self.assertIn(point.replace("：", ""), text.replace("：", ""))
                for shape in page.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                    self.assertLessEqual(shape.top + shape.height, prs.slide_height)
            self.assertEqual(sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in prs.slides[2].shapes), 1)

    def test_mixed_language_wrap_preserves_text_and_detects_overflow(self):
        text = "比较 API gateway 与数据处理条件，避免遗漏 latency 口径。"
        wrapped, size, fits = fit_text(text, 3.1, 1.3, 24, 16)
        self.assertTrue(fits)
        self.assertEqual(wrapped.replace("\n", "").replace(" ", ""), text.replace(" ", ""))
        self.assertFalse(fit_text("长文" * 300, 2, 1, 24, 16)[2])
        wrapped = wrap_text("条件需要现场核实。", 16*9, 16)
        self.assertTrue(all(not line.startswith("。") for line in wrapped.split("\n")))

    def test_template_brand_artwork_remains_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "模板标题"
            slide.placeholders[1].text = "模板正文"
            image = BytesIO()
            Image.new("RGB", (60, 30), "#223344").save(image, format="PNG")
            image.seek(0)
            slide.shapes.add_picture(image, 0, 0)
            template = root / "template.pptx"
            prs.save(template)
            output = root / "result.pptx"
            create_deck(Project(id="template-check", slides=[example_project().slides[1]]), output, template)
            original_picture = next(s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
            final_picture = next(s for s in Presentation(output).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
            self.assertEqual(original_picture.image.blob, final_picture.image.blob)


if __name__ == "__main__":
    unittest.main()
