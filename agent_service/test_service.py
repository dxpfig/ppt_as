"""Integration acceptance tests; requires a running service and real LibreOffice."""
import asyncio
from io import BytesIO
import os
from pathlib import Path
import time
import unittest
from uuid import uuid4
import zipfile

import httpx
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from agent_service.rendering import validate_package

BASE = os.getenv("STUDIO_TEST_URL", "http://127.0.0.1:8001/api/studio")


def company_template():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("EAF0FA")
    slide.shapes.title.text = "公司标题"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "WenQuanYi Zen Hei"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    slide.shapes.placeholders[1].text = "公司正文"
    image = BytesIO()
    Image.new("RGB", (80, 30), "#AB1234").save(image, "PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(8), Inches(.2), width=Inches(1))
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


class ServiceAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = httpx.Client(timeout=20)
        cls.project = {"imageSource": "none", "id": "acceptance-" + uuid4().hex, "title": "公司模板验收", "slides": [{"layout": "auto", "section": "", "source": "", "imageIntent": "auto", "imageQuery": "", "imageDisabled": False, "image": None, "id": "one", "title": "新的业务结论", "subtitle": "确认模板样式保留", "points": ["真实生成的内容", "保留公司标志与字体"], "notes": "这是演讲备注。"}], "theme": "forest", "phase": "ready", "messages": [], "versions": []}

    @classmethod
    def tearDownClass(cls):
        cls.client.close()

    def wait_job(self, identifier):
        for _ in range(160):
            job = self.client.get(f"{BASE}/jobs/{identifier}").json()
            if job["status"] in ("completed", "failed", "cancelled"):
                return job
            time.sleep(1)
        self.fail("render timed out")

    def test_projects_persist(self):
        url = f"{BASE}/projects/{self.project['id']}"
        self.assertEqual(self.client.put(url, json=self.project).status_code, 200)
        self.assertEqual(self.client.get(url).json()["slides"], self.project["slides"])
        self.assertTrue(any(p["id"] == self.project["id"] for p in self.client.get(f"{BASE}/projects").json()))

    def test_bad_upload_and_limits(self):
        response = self.client.post(f"{BASE}/templates", files={"file": ("broken.pptx", b"not a zip")})
        self.assertEqual(response.status_code, 400)
        response = self.client.post(f"{BASE}/templates", files={"file": ("macro.pptm", b"not a zip")})
        self.assertEqual(response.status_code, 400)
        with self.assertRaises(ValueError):
            validate_package(b"x" * (30 * 1024 * 1024 + 1))
        malicious = BytesIO()
        with zipfile.ZipFile(malicious, "w") as z:
            z.writestr("ppt/presentation.xml", '<!DOCTYPE x [<!ENTITY bomb "x">]><x/>')
        with self.assertRaises(ValueError):
            validate_package(malicious.getvalue())
        self.assertEqual(self.client.post(f"{BASE}/jobs", json={"kind": "render", "project": {**self.project, "slides": self.project["slides"] * 61}}).status_code, 422)

    def test_real_render(self):
        started = self.client.post(f"{BASE}/jobs", json={"kind": "render", "project": self.project})
        self.assertEqual(started.status_code, 202)
        job = self.wait_job(started.json()["id"])
        self.assertEqual(job["status"], "completed", job)
        self.check_artifact(job["result"]["artifact"])

    def check_artifact(self, artifact):
        origin = BASE.split("/api/studio")[0]
        self.assertEqual(len(artifact["previewUrls"]), 1)
        pptx = self.client.get(origin + artifact["pptxUrl"]).content
        self.assertEqual(Presentation(BytesIO(pptx)).slides[0].shapes.title.text if Presentation(BytesIO(pptx)).slides[0].shapes.title is not None else self.project["slides"][0]["title"], self.project["slides"][0]["title"])
        self.assertTrue(self.client.get(origin + artifact["pdfUrl"]).content.startswith(b"%PDF"))
        self.assertTrue(self.client.get(origin + artifact["previewUrls"][0]).content.startswith(b"\x89PNG"))
        return pptx

    def test_template_and_potx_preserve_styles(self):
        source = company_template()
        # Keep a real fixture for browser acceptance, no credentials or model data.
        fixture = Path("runtime/company-template-test.pptx")
        fixture.parent.mkdir(exist_ok=True)
        fixture.write_bytes(source)
        upload = self.client.post(f"{BASE}/templates", files={"file": ("公司模板.pptx", source)} )
        self.assertEqual(upload.status_code, 200, upload.text)
        template = upload.json()
        self.assertEqual(self.client.get(f"{BASE}/templates/{template['id']}").status_code, 200)
        started = self.client.post(f"{BASE}/jobs", json={"kind": "render", "project": {**self.project, "templateId": template["id"]}}).json()
        job = self.wait_job(started["id"])
        self.assertEqual(job["status"], "completed", job)
        pptx = self.check_artifact(job["result"]["artifact"])
        generated = Presentation(BytesIO(pptx))
        slide_text = "\n".join(shape.text for shape in generated.slides[0].shapes if shape.has_text_frame)
        for point in self.project["slides"][0]["points"]:
            self.assertIn(point, slide_text)
        self.assertEqual(generated.slide_width, Presentation(BytesIO(source)).slide_width)
        self.assertEqual(generated.slides[0].background.fill.fore_color.rgb, RGBColor.from_string("EAF0FA"))
        self.assertEqual(generated.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.name, "WenQuanYi Zen Hei")
        with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(BytesIO(pptx)) as final:
            media = [name for name in original.namelist() if name.startswith("ppt/media/")]
            self.assertTrue(media)
            for name in media:
                self.assertEqual(original.read(name), final.read(name))
            self.assertEqual(original.read("ppt/theme/theme1.xml"), final.read("ppt/theme/theme1.xml"))
        potx = BytesIO()
        with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(potx, "w") as converted:
            for name in original.namelist():
                content = original.read(name)
                if name == "[Content_Types].xml":
                    content = content.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
                converted.writestr(name, content)
        self.assertEqual(self.client.post(f"{BASE}/templates", files={"file": ("公司模板.potx", potx.getvalue())}).status_code, 200)

    def test_cancel_and_failure(self):
        started = self.client.post(f"{BASE}/jobs", json={"kind": "render", "project": self.project}).json()
        self.assertEqual(self.client.post(f"{BASE}/jobs/{started['id']}/cancel").json()["status"], "cancelled")
        self.assertEqual(self.wait_job(started["id"])["status"], "cancelled")
        failed = self.client.post(f"{BASE}/jobs", json={"kind": "render", "project": {**self.project, "slides": []}}).json()
        self.assertEqual(self.wait_job(failed["id"])["status"], "failed")


if __name__ == "__main__":
    unittest.main(verbosity=2)
