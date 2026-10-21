"""Generate the deterministic, sanitized company template used by golden tests."""
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'tests' / 'golden' / 'inputs' / 'company-template.pptx'


def main() -> None:
    presentation = Presentation()
    presentation.core_properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    presentation.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    presentation.core_properties.last_modified_by = 'Golden Fixture'
    presentation.core_properties.revision = 1
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    accent = RGBColor.from_string('1F5A7A')
    background = RGBColor.from_string('F3F6F8')

    for index, (title, body) in enumerate(
        (
            ('公司演示模板', '封面副标题'),
            ('内容页标题', '这里填写正文内容'),
            ('感谢观看', '模板结束页'),
        )
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = background
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8.5), Inches(1.0))
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.runs[0].font.name = 'WenQuanYi Zen Hei'
        title_paragraph.runs[0].font.size = Pt(30)
        title_paragraph.runs[0].font.color.rgb = accent
        body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.0), Inches(2.0))
        body_paragraph = body_box.text_frame.paragraphs[0]
        body_paragraph.text = body
        body_paragraph.runs[0].font.name = 'WenQuanYi Zen Hei'
        body_paragraph.runs[0].font.size = Pt(18)
        if index == 1:
            block = slide.shapes.add_shape(
                1,
                Inches(9.2),
                Inches(1.2),
                Inches(2.8),
                Inches(1.6),
            )
            block.fill.solid()
            block.fill.fore_color.rgb = accent
            block.line.fill.background()
        if index in (0, 2):
            image = BytesIO()
            Image.new('RGB', (120, 36), '#1F5A7A').save(image, 'PNG')
            image.seek(0)
            slide.shapes.add_picture(image, Inches(10.7), Inches(0.25), width=Inches(1.7))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    raw = BytesIO()
    presentation.save(raw)
    raw.seek(0)
    temporary = OUTPUT.with_suffix('.tmp.pptx')
    with ZipFile(raw) as source, ZipFile(
        temporary,
        'w',
        compression=ZIP_DEFLATED,
    ) as destination:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            destination.writestr(entry, source.read(name))
    temporary.replace(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
