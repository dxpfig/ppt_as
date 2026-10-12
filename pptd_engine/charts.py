"""PPTD chart export: native OOXML charts, or explicitly opted-in editable vectors.

The browser's validator is the single source of chart interpretation. Node/ECharts
is also used for SVG layout; SVG is converted to DrawingML shapes, never embedded
as an image. Original authored data is retained in a related custom XML part.
"""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
import io
import json
import math
from pathlib import Path
import re
import shutil
import subprocess

from lxml import etree
import xlsxwriter
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

PALETTE = ['#1783FE', '#67B8A6', '#C9A869', '#8A84C7', '#8C98A8', '#5E9CF4', '#7AB5D7', '#B08A9C', '#333333', '#DCDCDC']
NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RUNTIME = Path(__file__).resolve().parents[1] / 'packages/pptd-runtime/src/charts.js'
CUSTOM_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml'


def _node(element, theme, *, svg=False):
    node = shutil.which('node')
    if not node:
        raise ValueError('ChartRuntimeUnavailable: Node.js is required for the shared PPTD chart validator')
    script = "import {readFileSync} from 'node:fs'; const m=await import(process.argv[1]); const {element,theme,svg}=JSON.parse(readFileSync(0,'utf8')); try {const result=m.validateChart(element); m.buildChartOption(element,theme); if(svg) result.svg=m.renderChartSvg(element,theme); process.stdout.write(JSON.stringify(result));} catch(e) {process.stderr.write(e.message);process.exitCode=1;}"
    result = subprocess.run([node, '--input-type=module', '-e', script, RUNTIME.as_uri()], input=json.dumps({'element': element, 'theme': theme, 'svg': svg}), text=True, capture_output=True, timeout=60)
    if result.returncode:
        raise ValueError(result.stderr.strip() or 'ChartRuntimeError')
    return json.loads(result.stdout)


def _el(tag, parent=None, text=None, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(k, str(v))
    if text is not None:
        el.text = str(text)
    if parent is not None:
        parent.append(el)
    return el


def _color(value, theme):
    seen = set()
    while isinstance(value, str) and value.startswith('$'):
        if value in seen:
            raise ValueError('CyclicColorError: cyclic theme color reference')
        seen.add(value)
        value = theme.get('colors', {}).get(value[1:])
        if value is None:
            raise ValueError('UnknownColorError: unresolved theme color')
    if value in (None, 'none', 'transparent'):
        return '000000', 0
    if value.startswith('#'):
        h = value[1:]
        if len(h) == 3:
            h = ''.join(c * 2 for c in h)
        return h[:6].upper(), int(h[6:8], 16) / 255 if len(h) == 8 else 1
    match = re.match(r'rgba?\(([^)]+)\)', value)
    if match:
        parts = re.split(r'[, ]+', match[1])
        return ''.join(f'{round(float(n)):02X}' for n in parts[:3]), float(parts[3]) if len(parts) > 3 else 1
    named = {'white': 'FFFFFF', 'black': '000000', 'red': 'FF0000', 'blue': '0000FF', 'green': '008000', 'gray': '808080'}
    if value in named:
        return named[value], 1
    raise ValueError('InvalidValueError: unrecognized chart color ' + str(value))


def _solid(parent, color, theme, alpha=1):
    rgb, opacity = _color(color, theme)
    fill = _el('a:solidFill', parent)
    color_el = _el('a:srgbClr', fill, val=rgb)
    if opacity * alpha < 1:
        _el('a:alpha', color_el, val=round(max(0, opacity * alpha) * 100000))
    return fill


def _fill(parent, fill, theme, alpha=1):
    if fill is None or fill == 'none':
        return _el('a:noFill', parent)
    if isinstance(fill, str):
        return _solid(parent, fill, theme, alpha)
    if fill.get('type') == 'solid':
        return _solid(parent, fill['color'], theme, alpha)
    gradient = _el('a:gradFill', parent, rotWithShape='1')
    stops = _el('a:gsLst', gradient)
    for stop in fill['stops']:
        gs = _el('a:gs', stops, pos=round(stop['position'] * 100000))
        rgb, opacity = _color(stop['color'], theme)
        rgb_el = _el('a:srgbClr', gs, val=rgb)
        if opacity * alpha < 1:
            _el('a:alpha', rgb_el, val=round(opacity * alpha * 100000))
    if fill.get('gradientType') == 'radial':
        path = _el('a:path', gradient, path='circle')
        _el('a:fillToRect', path, l=50000, t=50000, r=50000, b=50000)
    else:
        _el('a:lin', gradient, ang=round(fill.get('angle', 0) * 60000), scaled='1')
    return gradient


def _line(parent, spec, theme, default=None):
    spec = spec if isinstance(spec, dict) else {}
    line = _el('a:ln', parent, w=round(spec.get('width', 1) * 12700))
    if not spec and default is None:
        _el('a:noFill', line)
    else:
        _fill(line, spec.get('color', default or '#000000'), theme)
        _el('a:prstDash', line, val={'dash': 'dash', 'dot': 'sysDot'}.get(spec.get('style'), 'solid'))
    return line


def _shape_style(parent, fill, line, theme, *, alpha=1):
    sp = _el('c:spPr', parent)
    _fill(sp, fill, theme, alpha)
    _line(sp, line, theme)
    return sp


def _shadow(parent, spec, theme):
    effects = _el('a:effectLst', parent)
    if spec:
        x, y = spec.get('offset', [0, 0])
        node = _el('a:outerShdw', effects, blurRad=round(spec.get('blur', 0) * 12700), dist=round(math.hypot(x, y) * 12700), dir=round(math.degrees(math.atan2(y, x)) % 360 * 60000), algn='ctr', rotWithShape=0)
        rgb, alpha = _color(spec.get('color', '#00000080'), theme)
        color = _el('a:srgbClr', node, val=rgb)
        if alpha < 1:
            _el('a:alpha', color, val=round(alpha * 100000))
    return effects


def _image_fill(owner_part, fill, theme, assets, bounds):
    from types import SimpleNamespace
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptd_engine.icons import _paint
    def add_image(filename):
        image_part = owner_part.package.get_or_add_image_part(filename)
        return image_part, owner_part.relate_to(image_part, RT.IMAGE)
    proxy = SimpleNamespace(part=SimpleNamespace(get_or_add_image_part=add_image))
    return _paint(proxy, fill, theme, 1, assets, bounds)


def _text_style(parent, style, chart, theme, tag='c:txPr', text=None):
    tx = _el(tag, parent)
    _el('a:bodyPr', tx)
    _el('a:lstStyle', tx)
    p = _el('a:p', tx)
    pp = _el('a:pPr', p)
    family = style.get('fontFamily') or chart.get('fontFamily') or theme.get('fontFamily') or theme.get('textStyles', {}).get('body', {}).get('fontFamily') or 'MiSans'
    latin = family if isinstance(family, str) else family['latin']
    ea = family if isinstance(family, str) else family['ea']
    size = style.get('fontSize', max(10, min(16, chart.get('bounds', [0, 0, 640, 360])[3] / 30)))
    def runstyle(tag, parent):
        rp = _el(tag, parent, sz=round(size * 100))
        _solid(rp, style.get('color', theme.get('colors', {}).get('text', '#000000')), theme)
        _el('a:latin', rp, typeface=latin)
        _el('a:ea', rp, typeface=ea)
    runstyle('a:defRPr', pp)
    if text is not None:
        r = _el('a:r', p)
        runstyle('a:rPr', r)
        _el('a:t', r, text=text)
    _el('a:endParaRPr', p)
    return tx


def _title(parent, title, chart, theme):
    if not title:
        return
    style = {'text': title} if isinstance(title, str) else title
    el = _el('c:title', parent)
    tx = _el('c:tx', el)
    _text_style(tx, style, chart, theme, 'c:rich', style['text'])
    _el('c:overlay', el, val=0)


def _labels(parent, s, chart, theme):
    labels = s.get('dataLabels', {})
    if not labels.get('show'):
        return
    content = labels.get('content', 'value')
    node = _el('c:dLbls', parent)
    if labels.get('numberFormat'):
        _el('c:numFmt', node, formatCode=labels['numberFormat'], sourceLinked=0)
    _text_style(node, labels, chart, theme)
    if labels.get('position'):
        _el('c:dLblPos', node, val={'top': 't', 'bottom': 'b', 'left': 'l', 'right': 'r',
            'inside': 'ctr', 'insideEnd': 'inEnd', 'outsideEnd': 'outEnd'}[labels['position']])
    for key, show in [('showLegendKey', False), ('showVal', content == 'value'), ('showCatName', content == 'category'), ('showSerName', False), ('showPercent', content == 'percentage')]:
        _el('c:' + key, node, val=int(show))
    return node


def _marker(parent, value, color, theme, *, required=False):
    spec = value if isinstance(value, dict) else {}
    marker = _el('c:marker', parent)
    symbol = 'none' if not value and not required else {'rect': 'square'}.get(spec.get('shape'), spec.get('shape', 'circle'))
    _el('c:symbol', marker, val=symbol)
    if symbol != 'none':
        _el('c:size', marker, val=max(2, min(72, round(spec.get('size', 7)))))
        _shape_style(marker, spec.get('fill', color), spec.get('border'), theme)


def _attach_source(slide, shape, element, mode):
    package = slide.part.package
    root = etree.Element('{urn:pptd:chart-source:v2}chart', nsmap={'pptd': 'urn:pptd:chart-source:v2'}, elementId=element.get('elementId', ''), shapeId=str(shape.shape_id), mode=mode)
    root.text = json.dumps(element, ensure_ascii=False, separators=(',', ':'))
    part = Part(package.next_partname('/customXml/pptdChart%d.xml'), 'application/xml', package, etree.tostring(root, xml_declaration=True, encoding='UTF-8'))
    rid = slide.part.relate_to(part, CUSTOM_REL)
    return {'part': str(part.partname), 'relationshipId': rid}


class _Workbook:
    def __init__(self, element):
        self.output = io.BytesIO()
        self.book = xlsxwriter.Workbook(self.output, {'in_memory': True})
        self.sheet = self.book.add_worksheet('Chart Data')
        self.col = 0
        source = self.book.add_worksheet('PPTD Source')
        for j, name in enumerate(element['data']['cols']):
            source.write(0, j, name)
        for i, row in enumerate(element['data']['rows']):
            for j, value in enumerate(row):
                if isinstance(value, str):
                    source.write_string(i + 1, j, value)
                elif value is not None:
                    source.write(i + 1, j, value)

    def ref(self, parent, tag, values, name, numeric=True):
        col = self.col
        self.col += 1
        self.sheet.write_string(0, col, name)
        for row, value in enumerate(values, 1):
            if value is not None:
                if numeric:
                    self.sheet.write_number(row, col, float(value))
                else:
                    self.sheet.write_string(row, col, str(value))
        node = _el('c:' + tag, parent)
        ref = _el('c:numRef' if numeric else 'c:strRef', node)
        letter = xlsxwriter.utility.xl_col_to_name(col)
        _el('c:f', ref, text=f"'Chart Data'!${letter}$2:${letter}${max(2, len(values) + 1)}")
        cache = _el('c:numCache' if numeric else 'c:strCache', ref)
        if numeric:
            _el('c:formatCode', cache, text='General')
        _el('c:ptCount', cache, val=len(values))
        for i, value in enumerate(values):
            if value is not None:
                pt = _el('c:pt', cache, idx=i)
                _el('c:v', pt, text=value if numeric else str(value))

    def finish(self):
        self.book.close()
        return self.output.getvalue()


def _native_gap(element, normalized, series):
    peers = [s for s in normalized['series'] if s['type'] == series['type'] and s.get('xAxisIndex', 0) == series.get('xAxisIndex', 0) and s.get('yAxisIndex', 0) == series.get('yAxisIndex', 0)]
    slots = (1 if any(s.get('stack') for s in peers) else 0) + sum(not s.get('stack') for s in peers)
    cluster = slots + max(0, slots - 1) * element.get('barGap', 0)
    gap = element.get('categoryGap', .2)
    return gap / (1 - gap) * cluster * 100


def _native_reason(normalized, element=None):
    if element and element.get('barWidth') is not None:
        return 'explicit category-slot bar width requires vector layout'
    series = normalized['series']
    bars = [s for s in series if s['type'] == 'bar']
    if any(s.get('stack') for s in bars) and any(not s.get('stack') for s in bars):
        return 'a shared stacked and unstacked bar cluster requires vector layout'
    if any(s['type'] in ['bar', 'candlestick'] and _native_gap(element or {}, normalized, s) > 500 for s in series):
        return 'category gap exceeds Office native 500% of bar width limit'
    for s in series:
        if s['type'] == 'area' and (s.get('smooth') or s.get('marker')):
            return 'area smoothing and markers require vector paths'
        if s['type'] in ['waterfall', 'heatmap', 'treemap', 'sunburst', 'sankey']:
            return s['type'] + ' has no equivalent legacy DrawingML chart object'
        if s['type'] in ['bar', 'line', 'area'] and (s['yType'] if s['horizontal'] else s['xType']) == 'value':
            return 'continuous category-coordinate spacing needs vector geometry'
        if s['type'] == 'pie' and 0 < s.get('innerRadius', 0) < .1 or s['type'] == 'pie' and s.get('innerRadius', 0) > .9:
            return 'requested donut radius lies outside Office native 10–90% range'
        if isinstance(s.get('marker'), dict) and not 2 <= s['marker'].get('size', 7) <= 72:
            return 'marker size lies outside Office native 2–72pt range'
        if s.get('stack') == 'stream':
            return 'streamgraph requires a centered custom baseline'
        if s.get('symbol'):
            return 'pictographic bars require custom vector geometry'
        if s['type'] == 'bubble' and (s.get('sizeScale', 'sqrt') != 'sqrt' or s.get('sizeRange')):
            return 'explicit bubble radius mapping is not represented by native Office size semantics'
        if s['type'] == 'radar' and (s.get('smooth') or any(p['y'] is None for p in s['points'])):
            return 'radar smoothing and missing-spoke policy require explicit vector paths'
        if s['type'] == 'candlestick' and normalized['types'].count('candlestick') and sum(t['type'] == 'candlestick' for t in series) > 1:
            return 'multiple independent stock bodies require vector composition'
    if any(s['type'] in ['scatter', 'bubble'] for s in series) and any(s['type'] in ['bar', 'line', 'area'] and s['xType'] == 'category' for s in series):
        return 'numeric scatter and category plots use incompatible native horizontal axes'
    return None


def _native(slide, element, theme, normalized, asset_index=None):
    data = CategoryChartData()
    data.categories = ['']
    data.add_series('PPTD', [0])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *[Pt(v) for v in element['bounds']], data)
    chart = shape.chart
    old_external = deepcopy(chart._chartSpace.find('{%s}externalData' % NS['c']))
    space = chart._chartSpace
    for child in list(space):
        space.remove(child)
    _el('c:date1904', space, val=0)
    _el('c:lang', space, val='en-US')
    chart_xml = _el('c:chart', space)
    _title(chart_xml, element.get('title'), element, theme)
    _el('c:autoTitleDeleted', chart_xml, val=int(not element.get('title')))
    plotarea = _el('c:plotArea', chart_xml)
    _el('c:layout', plotarea)
    workbook = _Workbook(element)
    groups = {}
    axes = {}
    next_series = 0
    type_index = defaultdict(int)
    for s in normalized['series']:
        typ = s['type']
        color = s.get('fill') or s.get('lineColor') or PALETTE[type_index[typ] % 10]
        type_index[typ] += 1
        if isinstance(color, list):
            color = color[0]
        xi, yi = s.get('xAxisIndex', 0), s.get('yAxisIndex', 0)
        x_id, y_id = 1000 + xi, 2000 + yi
        if typ == 'radar':
            x_type, y_type = 'category', 'value'
        else:
            x_type, y_type = s['xType'], s['yType']
        tag = {'bar': 'barChart', 'line': 'lineChart', 'area': 'areaChart', 'scatter': 'scatterChart', 'bubble': 'bubbleChart', 'radar': 'radarChart', 'candlestick': 'stockChart', 'pie': 'doughnutChart' if s.get('innerRadius', 0) > 0 else 'pieChart'}[typ]
        key = (tag, xi, yi, s.get('stack'), s['horizontal'])
        if key not in groups:
            node = _el('c:' + tag, plotarea)
            groups[key] = (node, s)
            if typ == 'bar':
                _el('c:barDir', node, val='bar' if s['horizontal'] else 'col')
            if typ in ['bar', 'line', 'area']:
                _el('c:grouping', node, val={'percent': 'percentStacked', 'value': 'stacked'}.get(s.get('stack'), 'clustered' if typ == 'bar' else 'standard'))
            if typ == 'scatter':
                _el('c:scatterStyle', node, val='marker')
            if typ == 'radar':
                _el('c:radarStyle', node, val='filled')
            if typ != 'candlestick':
                _el('c:varyColors', node, val=int(typ == 'pie'))
        node = groups[key][0]
        if typ not in ['pie']:
            axes[('x', xi)] = {'id': x_id, 'cross': y_id, 'type': x_type, 'radar': typ == 'radar', 'horizontal': s['horizontal']}
            axes[('y', yi)] = {'id': y_id, 'cross': x_id, 'type': y_type, 'radar': typ == 'radar', 'horizontal': s['horizontal']}
        channels = (['open'] if 'open' in s['encode'] else []) + ['high', 'low', 'close'] if typ == 'candlestick' else [None]
        for stock_channel in channels:
            ser = _el('c:ser', node)
            _el('c:idx', ser, val=next_series)
            _el('c:order', ser, val=next_series)
            next_series += 1
            tx = _el('c:tx', ser)
            _el('c:v', tx, text=s['encode'][stock_channel] if stock_channel else s['name'])
            stroke = s.get('border')
            fill = color
            alpha = 1
            if typ in ['line', 'radar', 'area']:
                stroke = {'color': s.get('lineColor', color), 'width': s.get('width', 2), 'style': s.get('lineStyle', 'solid')}
                fill = s.get('areaColor', color) if typ != 'line' else None
                alpha = 1 if s.get('areaColor') or typ == 'line' else .45
            if typ == 'scatter':
                fill, stroke = None, None
            if typ == 'candlestick':
                fill, stroke = None, None
            _shape_style(ser, fill, stroke, theme, alpha=alpha)
            if typ in ['line', 'radar', 'scatter']:
                marker = s.get('marker')
                if typ == 'scatter':
                    marker = {'border': s.get('border'), **(marker if isinstance(marker, dict) else {})}
                _marker(ser, marker, color, theme, required=typ == 'scatter')
            if typ == 'candlestick':
                _marker(ser, {'size': 5} if stock_channel == 'close' and 'open' not in s['encode'] else False, s.get('wickStyle', {}).get('color', color), theme)
            if typ == 'pie':
                fills = s.get('fill')
                for j, point in enumerate(s['points']):
                    dp = _el('c:dPt', ser)
                    _el('c:idx', dp, val=j)
                    fill = fills[j % len(fills)] if isinstance(fills, list) else fills or PALETTE[j % 10]
                    _shape_style(dp, fill, s.get('border'), theme)
            if typ != 'candlestick':
                _labels(ser, s, element, theme)
            points = s['points']
            if typ in ['scatter', 'bubble']:
                workbook.ref(ser, 'xVal', [p['x'] for p in points], s['encode']['x'])
                workbook.ref(ser, 'yVal', [p['y'] for p in points], s['encode']['y'])
                if typ == 'bubble':
                    workbook.ref(ser, 'bubbleSize', [p['size'] for p in points], s['encode']['size'])
                    _el('c:bubble3D', ser, val=0)
            else:
                cat = 'category' if typ in ['pie', 'radar'] else 'y' if s['horizontal'] else 'x'
                val = stock_channel or ('value' if typ == 'pie' else 'x' if s['horizontal'] else 'y')
                workbook.ref(ser, 'cat', [p[cat] for p in points], s['encode'][cat], numeric=False)
                workbook.ref(ser, 'val', [p[val] for p in points], s['encode'][val])
            if typ in ['line', 'scatter']:
                _el('c:smooth', ser, val=int(s.get('smooth', False)))
    for key, (plot, s) in groups.items():
        typ = s['type']
        if typ == 'bar':
            _el('c:gapWidth', plot, val=round(_native_gap(element, normalized, s)))
            _el('c:overlap', plot, val=100 if s.get('stack') else -round(element.get('barGap', 0) * 100))
        if typ == 'line':
            _el('c:marker', plot, val=1)
        if typ == 'candlestick':
            lines = _el('c:hiLowLines', plot)
            _shape_style(lines, None, s.get('wickStyle', {'color': '#333333', 'width': 1}), theme)
            if 'open' in s['encode']:
                bars = _el('c:upDownBars', plot)
                _el('c:gapWidth', bars, val=round(_native_gap(element, normalized, s)))
                for name, fallback in [('upBars', '#67B8A6'), ('downBars', '#D96060')]:
                    b = _el('c:' + name, bars)
                    spec = s.get(name, {})
                    _shape_style(b, spec.get('fill', fallback), spec.get('border'), theme)
        if typ == 'pie':
            _el('c:firstSliceAng', plot, val=round(s.get('startAngle', 0) % 360))
            if s.get('innerRadius', 0) > 0:
                _el('c:holeSize', plot, val=max(10, min(90, round(s['innerRadius'] * 100))))
        if typ == 'bubble':
            _el('c:bubbleScale', plot, val=100)
            _el('c:showNegBubbles', plot, val=0)
            _el('c:sizeRepresents', plot, val='area')
        if typ != 'pie':
            _el('c:axId', plot, val=1000 + s.get('xAxisIndex', 0))
            _el('c:axId', plot, val=2000 + s.get('yAxisIndex', 0))
    for (side, index), info in axes.items():
        typ = info['type']
        raw = element.get(side + 'Axis', {})
        cfg = raw[index] if isinstance(raw, list) else raw
        if info['radar']:
            cfg = element.get('spokeAxis', {})
        axis = _el('c:catAx' if typ == 'category' else 'c:valAx', plotarea)
        _el('c:axId', axis, val=info['id'])
        scaling = _el('c:scaling', axis)
        _el('c:orientation', scaling, val='maxMin' if cfg.get('reverse') else 'minMax')
        if typ == 'value':
            for k in ['max', 'min']:
                if k in cfg:
                    _el('c:' + k, scaling, val=cfg[k])
        _el('c:delete', axis, val=int(not cfg.get('show', True)))
        _el('c:axPos', axis, val=('t' if index else 'b') if side == 'x' else ('r' if index else 'l'))
        grid = cfg.get('gridLine', True)
        if grid is not False:
            major = _el('c:majorGridlines', axis)
            _shape_style(major, None, grid if isinstance(grid, dict) else {'color': '#CCCCCC'}, theme)
        _title(axis, cfg.get('title'), element, theme)
        label = cfg.get('label', True)
        if isinstance(label, dict) and label.get('numberFormat'):
            _el('c:numFmt', axis, formatCode=label['numberFormat'], sourceLinked=0)
        _el('c:majorTickMark', axis, val='none')
        _el('c:minorTickMark', axis, val='none')
        _el('c:tickLblPos', axis, val='none' if label is False else 'nextTo')
        line_cfg = cfg.get('axisLine', True)
        sp = _shape_style(axis, None, line_cfg if isinstance(line_cfg, dict) else {'color': '#888888'} if line_cfg else None, theme)
        if isinstance(line_cfg, dict) and line_cfg.get('arrow'):
            arrow = line_cfg['arrow']
            ln = sp.find('{%s}ln' % NS['a'])
            if arrow in ('start', 'both'):
                _el('a:headEnd', ln, type='triangle')
            if arrow in (True, 'end', 'both'):
                _el('a:tailEnd', ln, type='triangle')
        tx = _text_style(axis, label if isinstance(label, dict) else {}, element, theme)
        if isinstance(label, dict) and 'rotation' in label:
            tx.find('a:bodyPr', NS).set('rot', str(round(-label['rotation'] * 60000)))
        _el('c:crossAx', axis, val=info['cross'])
        _el('c:crosses', axis, val='max' if index else 'autoZero')
        if typ == 'category':
            _el('c:auto', axis, val=1)
            _el('c:lblAlgn', axis, val='ctr')
            _el('c:lblOffset', axis, val=100)
            if isinstance(label, dict) and 'interval' in label:
                _el('c:tickLblSkip', axis, val=label['interval'] + 1)
        else:
            _el('c:crossBetween', axis, val='between')
    legend = element.get('legend', True)
    if legend is not False and (not isinstance(legend, dict) or legend.get('show', True)):
        cfg = legend if isinstance(legend, dict) else {}
        el = _el('c:legend', chart_xml)
        _el('c:legendPos', el, val={'top': 't', 'bottom': 'b', 'left': 'l', 'right': 'r'}.get(cfg.get('position', 'bottom'), 'b'))
        _el('c:overlay', el, val=0)
        _text_style(el, cfg, element, theme)
    _el('c:plotVisOnly', chart_xml, val=1)
    _el('c:dispBlanksAs', chart_xml, val={'gap': 'gap', 'connect': 'span', 'zero': 'zero'}[normalized['nullHandling']])
    frame_fill = element.get('fill')
    frame = _shape_style(space, None if isinstance(frame_fill, dict) and frame_fill.get('type') == 'image' else frame_fill, element.get('border'), theme)
    if isinstance(frame_fill, dict) and frame_fill.get('type') == 'image':
        frame.remove(frame.find('{%s}noFill' % NS['a']))
        frame.insert(0, _image_fill(chart.part, frame_fill, theme, asset_index, element['bounds']))
    _shadow(frame, element.get('shadow'), theme)
    _text_style(space, {}, element, theme)
    if old_external is not None:
        space.append(old_external)
    chart.part.chart_workbook.update_from_xlsx_blob(workbook.finish())
    return shape


def _svg_style(node, inherited):
    style = dict(inherited)
    style.update({k: v for k, v in node.attrib.items() if k in ['fill', 'stroke', 'opacity', 'fill-opacity', 'stroke-opacity', 'stroke-width', 'stroke-dasharray', 'font-size', 'font-family', 'font-weight', 'font-style', 'text-anchor', 'dominant-baseline']})
    for entry in node.get('style', '').split(';'):
        if ':' in entry:
            k, v = entry.split(':', 1)
            style[k.strip()] = v.strip()
            if k.strip() == 'font':
                match = re.fullmatch(r'(.*?)\s*(\d+(?:\.\d+)?)(px|pt)(?:/[^ ]+)?\s+(.+)', v.strip())
                if match:
                    prefix, size, unit, family = match.groups()
                    style.update({'font-size':size+unit,'font-family':family})
                    if 'bold' in prefix or '700' in prefix: style['font-weight']='bold'
                    if 'italic' in prefix: style['font-style']='italic'
    return style


def _svg_transform(value):
    from fontTools.misc.transform import Transform
    transform = Transform()
    for command, args in re.findall(r'(matrix|translate|scale|rotate|skewX|skewY)\s*\(([^)]*)\)', value or ''):
        n = [float(x) for x in re.findall(r'[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?', args)]
        if command == 'matrix':
            operation = Transform(*n)
        elif command == 'translate':
            operation = Transform().translate(n[0], n[1] if len(n) > 1 else 0)
        elif command == 'scale':
            operation = Transform().scale(n[0], n[1] if len(n) > 1 else n[0])
        elif command == 'rotate':
            operation = Transform()
            if len(n) == 3:
                operation = operation.translate(n[1], n[2])
            operation = operation.rotate(math.radians(n[0]))
            if len(n) == 3:
                operation = operation.translate(-n[1], -n[2])
        elif command == 'skewX':
            operation = Transform(1, 0, math.tan(math.radians(n[0])), 1, 0, 0)
        else:
            operation = Transform(1, math.tan(math.radians(n[0])), 0, 1, 0, 0)
        transform = transform.transform(operation)
    return transform


def _svg_path(node):
    tag = etree.QName(node).localname
    f = lambda k, default=0: float(node.get(k, default))
    if tag == 'path':
        return node.get('d', '')
    if tag == 'rect':
        x, y, w, h = f('x'), f('y'), f('width'), f('height')
        return f'M{x} {y}h{w}v{h}h{-w}Z'
    if tag in ('circle', 'ellipse'):
        x, y = f('cx'), f('cy')
        rx, ry = (f('r'), f('r')) if tag == 'circle' else (f('rx'), f('ry'))
        return f'M{x-rx} {y}A{rx} {ry} 0 1 0 {x+rx} {y}A{rx} {ry} 0 1 0 {x-rx} {y}Z'
    if tag == 'line':
        return f'M{f("x1")} {f("y1")}L{f("x2")} {f("y2")}'
    if tag in ('polyline', 'polygon'):
        return 'M' + node.get('points', '') + ('Z' if tag == 'polygon' else '')
    return ''


def _clip_path(data, rectangle, filled):
    """Clip SVG geometry to an ECharts rectangular viewport.

    Curves entirely inside the viewport remain curves. Crossing curves are
    flattened adaptively to <=0.05px error before polygon/segment clipping.
    """
    from fontTools.pens.basePen import BasePen
    from fontTools.pens.boundsPen import BoundsPen
    from fontTools.svgLib.path import parse_path
    bounds = BoundsPen(None)
    parse_path(data, bounds)
    if not bounds.bounds:
        return ''
    x0, y0, x1, y1 = rectangle
    bx0, by0, bx1, by1 = bounds.bounds
    if bx0 >= x0 and by0 >= y0 and bx1 <= x1 and by1 <= y1:
        return data
    class FlatPen(BasePen):
        def __init__(self):
            super().__init__(None)
            self.paths = []
            self.points = []
        def _moveTo(self, pt):
            self.points = [pt]
        def _lineTo(self, pt):
            self.points.append(pt)
        def _curveToOne(self, a, b, c):
            start = self._getCurrentPoint()
            def flatten(p0, p1, p2, p3, depth=0):
                dx, dy = p3[0]-p0[0], p3[1]-p0[1]
                den = math.hypot(dx, dy)
                error = max(abs(dy*(p1[0]-p0[0])-dx*(p1[1]-p0[1])), abs(dy*(p2[0]-p0[0])-dx*(p2[1]-p0[1]))) / max(den, .001)
                if error <= .05 or depth >= 12:
                    self.points.append(p3)
                    return
                mid = lambda p,q: ((p[0]+q[0])/2, (p[1]+q[1])/2)
                a,b,c = mid(p0,p1),mid(p1,p2),mid(p2,p3)
                d,e = mid(a,b),mid(b,c)
                f = mid(d,e)
                flatten(p0,a,d,f,depth+1)
                flatten(f,e,c,p3,depth+1)
            flatten(start,a,b,c)
        def _closePath(self):
            self.paths.append((self.points, True))
            self.points = []
        def _endPath(self):
            self.paths.append((self.points, False))
            self.points = []
    pen = FlatPen()
    parse_path(data, pen)
    if pen.points:
        pen.paths.append((pen.points, False))
    out = []
    for points, closed in pen.paths:
        if filled or closed:
            for axis, bound, lower in [(0,x0,True),(0,x1,False),(1,y0,True),(1,y1,False)]:
                if not points:
                    break
                output = []
                previous = points[-1]
                inside = lambda p: p[axis] >= bound if lower else p[axis] <= bound
                for current in points:
                    pi, ci = inside(previous), inside(current)
                    if pi != ci:
                        t = (bound-previous[axis])/(current[axis]-previous[axis])
                        output.append((previous[0]+t*(current[0]-previous[0]), previous[1]+t*(current[1]-previous[1])))
                    if ci:
                        output.append(current)
                    previous = current
                points = output
            if points:
                out.append('M' + 'L'.join(f'{x} {y}' for x,y in points) + 'Z')
        else:
            for a,b in zip(points, points[1:]):
                dx,dy=b[0]-a[0],b[1]-a[1]
                low,high=0,1
                accepted=True
                for p,q in [(-dx,a[0]-x0),(dx,x1-a[0]),(-dy,a[1]-y0),(dy,y1-a[1])]:
                    if p==0:
                        if q<0: accepted=False;break
                    elif p<0: low=max(low,q/p)
                    else: high=min(high,q/p)
                if accepted and low<=high:
                    out.append(f'M{a[0]+low*dx} {a[1]+low*dy}L{a[0]+high*dx} {a[1]+high*dy}')
    return ''.join(out)


def _svg_to_vectors(slide, svg, bounds, theme, element=None, asset_index=None):
    from fontTools.misc.transform import Transform
    from fontTools.pens.svgPathPen import SVGPathPen
    from fontTools.pens.boundsPen import BoundsPen
    from fontTools.pens.transformPen import TransformPen
    from fontTools.svgLib.path import parse_path
    from pptd_engine.icons import svg_path_geometry

    root = etree.fromstring(svg.encode() if isinstance(svg, str) else svg, etree.XMLParser(resolve_entities=False, no_network=True))
    width, height = float(root.get('width', bounds[2])), float(root.get('height', bounds[3]))
    group = slide.shapes.add_group_shape()
    frame = group.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Pt(width), Pt(height))
    frame.shadow.inherit = False
    for effect in frame._element.xpath('p:style/a:effectRef'):
        effect.set('idx', '0')
    frame.fill.background()
    frame.line.fill.background()
    element = element or {}
    if element.get('fill', {}).get('type') == 'image':
        sp = frame._element.spPr
        for child in list(sp):
            if etree.QName(child).localname in ['noFill', 'solidFill', 'gradFill', 'blipFill']:
                sp.remove(child)
        sp.insert(2, _image_fill(slide.part, element['fill'], theme, asset_index, bounds))
    if element.get('border'):
        sp = frame._element.spPr
        ln = sp.find('{%s}ln' % NS['a'])
        if ln is not None: sp.remove(ln)
        _line(sp, element['border'], theme)
    if element.get('shadow'):
        sp = frame._element.spPr
        effects = sp.find('{%s}effectLst' % NS['a'])
        if effects is not None: sp.remove(effects)
        _shadow(sp, element['shadow'], theme)
    gradients = {}
    for node in root.iter():
        tag = etree.QName(node).localname
        if tag not in ['linearGradient', 'radialGradient']:
            continue
        stops = []
        for child in node:
            if etree.QName(child).localname == 'stop':
                style = _svg_style(child, {})
                color = child.get('stop-color', style.get('stop-color', '#000000'))
                rgb, alpha = _color(color, theme)
                opacity = float(child.get('stop-opacity', style.get('stop-opacity', 1))) * alpha
                offset = child.get('offset', '0')
                stops.append({'position': float(offset.strip('%')) / (100 if '%' in offset else 1), 'color': '#' + rgb + (f'{round(opacity * 255):02X}' if opacity < 1 else '')})
        if tag == 'linearGradient':
            dx = float(node.get('x2', 1)) - float(node.get('x1', 0))
            dy = float(node.get('y2', 0)) - float(node.get('y1', 0))
            angle = math.degrees(math.atan2(dy, dx)) % 360
        else:
            angle = 0
        gradients[node.get('id')] = {'type': 'gradient', 'gradientType': 'radial' if tag == 'radialGradient' else 'linear', 'angle': angle, 'stops': stops}

    clips = {}
    for node in root.iter():
        if etree.QName(node).localname == 'clipPath':
            if node.get('clipPathUnits', 'userSpaceOnUse') != 'userSpaceOnUse' or len(node) != 1:
                raise ValueError('UnsupportedSvgClipError: expected a single viewport clip')
            child = node[0]
            pen = BoundsPen(None)
            parse_path(_svg_path(child), TransformPen(pen, _svg_transform(child.get('transform'))))
            clips[node.get('id')] = pen.bounds
    warnings = []
    def visit(node, transform=Transform(), inherited=None, opacity=1, clip=None):
        tag = etree.QName(node).localname
        if tag in ['defs', 'clipPath', 'style', 'metadata', 'title', 'desc']:
            return
        style = _svg_style(node, inherited or {})
        own_opacity = float(node.get('opacity', 1))
        transform = transform.transform(_svg_transform(node.get('transform')))
        opacity *= own_opacity
        if node.get('clip-path'):
            ref = node.get('clip-path').removeprefix('url(#').removesuffix(')')
            rect = clips.get(ref)
            if not rect:
                raise ValueError('SVGClipError: missing clip definition')
            p0 = transform.transformPoint((rect[0], rect[1]))
            p1 = transform.transformPoint((rect[2], rect[3]))
            new_clip = (min(p0[0],p1[0]), min(p0[1],p1[1]), max(p0[0],p1[0]), max(p0[1],p1[1]))
            clip = (max(clip[0],new_clip[0]), max(clip[1],new_clip[1]), min(clip[2],new_clip[2]), min(clip[3],new_clip[3])) if clip else new_clip
        if tag == 'text':
            text = ''.join(node.itertext())
            if not text:
                return
            font_size = float(re.sub(r'px|pt', '', style.get('font-size', '12')))
            # ECharts SVG text is already positioned; retain it as editable text.
            x, y = transform.transformPoint((float(node.get('x', 0)), float(node.get('y', 0))))
            text_width = max(font_size * 2, sum(font_size * (1 if ord(ch) > 255 else .8) for ch in text) + font_size)
            anchor = style.get('text-anchor', 'start')
            if anchor == 'middle':
                x -= text_width / 2
            elif anchor == 'end':
                x -= text_width
            baseline = style.get('dominant-baseline', 'alphabetic')
            y -= font_size * (.5 if baseline in ['central', 'middle'] else 0 if baseline in ['hanging', 'text-before-edge'] else .82)
            shape = group.shapes.add_textbox(Pt(x), Pt(y), Pt(text_width + 3), Pt(font_size * 1.4))
            shape.shadow.inherit = False
            shape.line.fill.background()
            shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            shape.text_frame.margin_left = shape.text_frame.margin_right = shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
            shape.text_frame.word_wrap = False
            shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.alignment = {'middle': PP_ALIGN.CENTER, 'end': PP_ALIGN.RIGHT}.get(anchor, PP_ALIGN.LEFT)
            run = paragraph.add_run()
            run.text = text
            run.font.name = style.get('font-family', 'MiSans').split(',')[0].strip('"\' ')
            run.font.size = Pt(font_size)
            run.font.bold = style.get('font-weight') in ['bold', '700', '800', '900']
            run.font.italic = style.get('font-style') == 'italic'
            rpr = run._r.get_or_add_rPr()
            solid = _solid(rpr, style.get('fill', '#000000'), theme, opacity * float(style.get('fill-opacity', 1)))
            rpr.remove(solid)
            rpr.insert(0, solid)
            _el('a:ea', rpr, typeface=run.font.name)
            angle = math.degrees(math.atan2(transform.xy, transform.xx))
            if angle:
                shape.rotation = angle
            return
        path = _svg_path(node)
        if path:
            pen = SVGPathPen(None)
            parse_path(path, TransformPen(pen, transform))
            commands = pen.getCommands()
            if clip:
                commands = _clip_path(commands, clip, style.get('fill', '#000000') != 'none')
            if not commands:
                return
            bbox = BoundsPen(None)
            parse_path(commands, bbox)
            if not bbox.bounds:
                return
            bx, by, bx1, by1 = bbox.bounds
            bw, bh = max(.01, bx1-bx), max(.01, by1-by)
            local = SVGPathPen(None)
            parse_path(commands, TransformPen(local, Transform().translate(-bx,-by)))
            shape = group.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(bx), Pt(by), Pt(bw), Pt(bh))
            shape.shadow.inherit = False
            for effect in shape._element.xpath('p:style/a:effectRef'):
                effect.set('idx', '0')
            sp = shape._element.spPr
            for child in list(sp):
                if etree.QName(child).localname != 'xfrm':
                    sp.remove(child)
            sp.append(svg_path_geometry([local.getCommands()], bw, bh))
            value = style.get('fill', '#000000')
            if value.startswith('url('):
                value = gradients.get(value.removeprefix('url(#').removesuffix(')'))
                if not value:
                    raise ValueError('SVGGradientError: missing gradient definition')
            if isinstance(value, dict) and value.get('gradientType') == 'linear':
                value = deepcopy(value)
                theta = math.radians(value.get('angle', 0))
                dx = math.cos(theta) * transform.xx + math.sin(theta) * transform.yx
                dy = math.cos(theta) * transform.xy + math.sin(theta) * transform.yy
                value['angle'] = math.degrees(math.atan2(dy, dx)) % 360
            _fill(sp, value, theme, opacity * float(style.get('fill-opacity', 1)))
            stroke = style.get('stroke', 'none') if float(style.get('stroke-width', 1)) > 0 else 'none'
            line = _el('a:ln', sp, w=round(float(style.get('stroke-width', 1)) * 12700))
            _fill(line, stroke, theme, opacity * float(style.get('stroke-opacity', 1)))
            if style.get('stroke-dasharray') not in [None, 'none']:
                dash = [float(x) for x in re.split('[, ]+', style['stroke-dasharray']) if x]
                custom = _el('a:custDash', line)
                for i in range(0, len(dash), 2):
                    _el('a:ds', custom, d=round(dash[i] * 100000), sp=round(dash[(i+1) % len(dash)] * 100000))
            _el('a:effectLst', sp)
        for child in node:
            visit(child, transform, style, opacity, clip)
    visit(root)
    group.left, group.top, group.width, group.height = [Pt(v) for v in bounds]
    group._element.grpSpPr.xfrm.chOff.x = 0
    group._element.grpSpPr.xfrm.chOff.y = 0
    group._element.grpSpPr.xfrm.chExt.cx = Pt(width)
    group._element.grpSpPr.xfrm.chExt.cy = Pt(height)
    return group, warnings


def add_chart(slide, element, theme=None, *, allow_vector=False, svg=None, asset_index=None):
    """Return ``(editable_shape_or_group, audit)``; reject implicit vector fallback."""
    theme = theme or {}
    normalized = _node(element, theme)
    reason = _native_reason(normalized, element)
    warnings = list(normalized['warnings'])
    if reason:
        if not allow_vector:
            raise ValueError('EditableVectorRequired: ' + reason + '; explicitly allow vector export or use strict native-compatible settings')
        if svg is None:
            visual = deepcopy(element)
            if visual.get('fill', {}).get('type') == 'image':
                visual.pop('fill')
            svg = _node(visual, theme, svg=True)['svg']
        shape, vector_warnings = _svg_to_vectors(slide, svg, element['bounds'], theme, element, asset_index)
        warnings.extend(vector_warnings)
        mode = 'editable-vector'
    else:
        shape = _native(slide, element, theme, normalized, asset_index)
        mode = 'native'
    shape.name = element.get('elementId', 'PPTD chart')
    source = _attach_source(slide, shape, element, mode)
    audit = {'mode': mode, 'types': normalized['types'], 'warnings': warnings, 'sourcePreserved': True, 'source': source}
    if reason:
        audit['reason'] = reason
    return shape, audit
