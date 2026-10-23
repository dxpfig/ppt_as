"""Behavior contracts for editable Office math and Font Awesome geometry."""
import io
import json
import unittest
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from pptd_engine.math import append_math, split_math, mathml_to_omml
from pptd_engine.fonts import canonical_family
from pptd_engine.icons import add_icon, load_icon_index, svg_path_geometry

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


class MathTests(unittest.TestCase):
    def paragraph(self):
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        return self.slide.shapes.add_textbox(0, 0, Pt(400), Pt(200)).text_frame.paragraphs[0]

    def test_structural_math_objects(self):
        cases = [
            (r"\frac{1}{3}", "f"), (r"\sqrt{x}", "rad"),
            (r"\sqrt[3]{x}", "rad"), (r"x_1^2", "sSubSup"),
            (r"\begin{pmatrix}a&b\\c&d\end{pmatrix}", "m"),
            (r"\int_0^1 x^2\mathrm{d}x", "nary"),
        ]
        for latex, tag in cases:
            with self.subTest(latex=latex):
                p = self.paragraph()
                append_math(p, latex, 24, "#123456")
                root = etree.fromstring(p._p.xml.encode())
                self.assertTrue(root.xpath(f"a14:m/m:oMath//m:{tag}", namespaces=NS))
                self.assertFalse(root.xpath(".//a:blip", namespaces=NS))

    def test_mixed_text_order_and_math_styles_survive_save(self):
        p = self.paragraph()
        before = p.add_run(); before.text = "中文结果："; before.font.bold = True
        append_math(p, r"\frac{a}{b}", 20, "#1278AB")
        after = p.add_run(); after.text = "，详见链接"; after.hyperlink.address = "https://example.org"
        output = io.BytesIO(); self.prs.save(output)
        with ZipFile(output) as z:
            root = etree.fromstring(z.read("ppt/slides/slide1.xml"))
        p = root.xpath(".//a:p", namespaces=NS)[0]
        self.assertEqual([etree.QName(e).localname for e in p], ["r", "m", "r"])
        self.assertEqual(p.xpath("a:r/a:t/text()", namespaces=NS), ["中文结果：", "，详见链接"])
        self.assertTrue(p.xpath("a:r/a:rPr/a:hlinkClick", namespaces=NS))
        self.assertTrue(p.xpath("a14:m//a:rPr[@sz='2000']/a:solidFill/a:srgbClr[@val='1278AB']", namespaces=NS))
        self.assertFalse(p.xpath("a14:m//a:rPr[@b='1'] | a14:m//a:hlinkClick", namespaces=NS))
        self.assertTrue(p.xpath("a14:m//a:latin[@typeface='STIX Two Math']", namespaces=NS))

    def test_math_in_table_and_chinese(self):
        self.paragraph()
        table = self.slide.shapes.add_table(1, 1, 0, 0, Pt(300), Pt(100)).table
        p = table.cell(0, 0).text_frame.paragraphs[0]
        append_math(p, r"\text{中文}+\sqrt{x}", 18, "#000000")
        self.assertIn("中文", p._p.xml)
        self.assertTrue(etree.fromstring(p._p.xml.encode()).xpath("a14:m/m:oMath/m:r/m:t", namespaces=NS))
        self.assertTrue(etree.fromstring(p._p.xml.encode()).xpath("a14:m//a:ea[@typeface='%s']" % canonical_family('Noto Sans CJK SC'), namespaces=NS))

    def test_bad_latex_and_unknown_mathml_are_explicit_errors(self):
        for latex in [r"\frac{a}", r"\unknowncommand{x}", r"{x", r"<b>x</b>"]:
            with self.subTest(latex=latex), self.assertRaises(ValueError):
                append_math(self.paragraph(), latex)
        with self.assertRaisesRegex(ValueError, "unsupported.*maction"):
            mathml_to_omml('<math xmlns="http://www.w3.org/1998/Math/MathML"><maction><mi>x</mi></maction></math>')

    def test_delimiters_preserve_plain_text_and_reject_unbalanced_math(self):
        self.assertEqual(split_math(r"中文 \(x^2\) 与 \(y\)"), [(False, "中文 "), (True, "x^2"), (False, " 与 "), (True, "y")])
        with self.assertRaisesRegex(ValueError, "delimiter"):
            split_math(r"broken \(x")


class IconTests(unittest.TestCase):
    def test_free_icon_catalog_covers_all_styles_and_aliases(self):
        index = load_icon_index()
        for name in ["fas:house", "fas:lightbulb", "far:heart", "far:calendar", "fab:github", "fab:python"]:
            self.assertGreater(index[name]["width"], 0)
            self.assertTrue(index[name]["paths"])
        self.assertGreater(len(index), 2000)

    def test_icon_is_editable_geometry_with_name_and_transform(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, name in enumerate(["fas:house", "far:heart", "fab:github"]):
            shape = add_icon(slide, {"elementId": f"icon-{i}", "iconName": name, "bounds": [5, 10, 32, 48], "rotation": 25, "flip": [True, False], "opacity": 0.5, "fill": {"type": "solid", "color": "$primary"}}, {"colors": {"primary": "#112233"}})
            root = etree.fromstring(shape._element.xml.encode())
            self.assertTrue(root.xpath(".//a:custGeom/a:pathLst/a:path", namespaces=NS))
            self.assertFalse(root.xpath(".//a:blip", namespaces=NS))
            self.assertTrue(root.xpath(".//a:xfrm[@flipH='1'][@rot='1500000']", namespaces=NS))
            self.assertIn(name, shape._element.xml)
            self.assertTrue(root.xpath(".//a:solidFill/a:srgbClr[@val='112233']/a:alpha[@val='50000']", namespaces=NS))

    def test_multipath_and_curves_are_preserved(self):
        g = svg_path_geometry(["M0 0L10 0L10 10Z", "M2 2C3 4 5 6 8 8Q9 9 2 2Z"], 10, 10)
        self.assertEqual(len(g.xpath("a:pathLst/a:path", namespaces=NS)), 2)
        self.assertTrue(g.xpath(".//a:cubicBezTo", namespaces=NS))
        self.assertTrue(g.xpath(".//a:quadBezTo", namespaces=NS))

    def test_unknown_icon_is_error(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaisesRegex(ValueError, "Unknown.*icon"):
            add_icon(slide, {"iconName": "fas:no-such-icon", "bounds": [0, 0, 30, 30]})

    def test_radial_gradient_remains_radial(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = add_icon(slide, {"iconName": "fas:heart", "bounds": [0, 0, 30, 30], "fill": {"type": "gradient", "gradientType": "radial", "stops": [{"position": 0, "color": "#ffffff"}, {"position": 1, "color": "#ff0000"}]}})
        root = etree.fromstring(shape._element.xml.encode())
        self.assertTrue(root.xpath(".//a:gradFill/a:path[@path='circle']", namespaces=NS))


if __name__ == "__main__":
    unittest.main()
