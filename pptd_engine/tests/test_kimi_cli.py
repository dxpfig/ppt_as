"""Unit coverage for the kimi-slides CLI parity surface."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import yaml

from pptd_engine import check as check_module
from pptd_engine import cli as cli_module
from pptd_engine.check import parse_page_spec
from pptd_engine.document import load_yaml


def _write_deck(root: Path, page_yaml: str) -> Path:
    (root / "pages").mkdir(exist_ok=True)
    (root / "deck.pptd").write_text(
        "version: v2\ntitle: test\nsize: [960, 540]\npages:\n- pages/page-1.page\n",
        encoding="utf-8",
    )
    (root / "pages" / "page-1.page").write_text(page_yaml, encoding="utf-8")
    return root / "deck.pptd"


class PageSpecTests(unittest.TestCase):
    def test_specs(self):
        self.assertEqual(parse_page_spec(None, 10), list(range(1, 11)))
        self.assertEqual(parse_page_spec("3", 10), [3])
        self.assertEqual(parse_page_spec("1,2", 10), [1, 2])
        self.assertEqual(parse_page_spec("2-10", 10), list(range(2, 11)))
        self.assertEqual(parse_page_spec("1,3-5", 10), [1, 3, 4, 5])
        with self.assertRaises(ValueError):
            parse_page_spec("0", 10)
        with self.assertRaises(ValueError):
            parse_page_spec("5-2", 10)


class CheckTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_issue_name_mapping_and_severity(self):
        _write_deck(
            self.root,
            "elements:\n- elementType: shape\n  shapeName: roundRect\n  bounds: [0, 0, 50, 50]\n",
        )
        report = check_module.check_deck(self.root, level="keep")
        self.assertTrue(any(i["name"] == "MissingField" and i["severity"] == "error" for i in report))

    def test_auto_repair_coerces_and_deletes(self):
        _write_deck(
            self.root,
            "elements:\n"
            "- elementId: t1\n"
            "  elementType: text\n"
            "  content:\n    text: hello\n    fontSize: '24'\n"
            "  bounds: [0, 0, 200, 200]\n"
            "- elementType: shape\n"
            "  shapeName: roundRect\n"
            "  bounds: [300, 300, 50, 50]\n",
        )
        report = check_module.check_deck(self.root, level="auto")
        fixed = [i for i in report if i["fixed"]]
        self.assertTrue(any(i["code"] == "InvalidType" and i["location"].endswith("content.fontSize") for i in fixed))
        self.assertTrue(any(i["code"] == "MissingField" and i["location"].endswith("elements[1]") for i in fixed))
        page = load_yaml(self.root / "pages" / "page-1.page")
        self.assertEqual(len(page["elements"]), 1)
        self.assertEqual(page["elements"][0]["content"]["fontSize"], 24)

    def test_keep_does_not_mutate(self):
        _write_deck(
            self.root,
            "elements:\n- elementType: shape\n  shapeName: roundRect\n  bounds: [0, 0, 50, 50]\n",
        )
        check_module.check_deck(self.root, level="keep")
        page = load_yaml(self.root / "pages" / "page-1.page")
        self.assertEqual(len(page["elements"]), 1)

    def test_page_spec_limits_structural_results_to_selected_page(self):
        (self.root / "pages").mkdir()
        (self.root / "deck.pptd").write_text(
            "version: v2\ntitle: test\nsize: [960, 540]\n"
            "pages:\n- pages/page-1.page\n- pages/page-2.page\n",
            encoding="utf-8",
        )
        (self.root / "pages" / "page-1.page").write_text(
            "elements:\n"
            "- elementId: ok\n"
            "  elementType: text\n"
            "  bounds: [20, 20, 200, 60]\n"
            "  content: {text: 第一页}\n",
            encoding="utf-8",
        )
        (self.root / "pages" / "page-2.page").write_text(
            "elements:\n"
            "- elementType: text\n"
            "  bounds: [20, 20, 200, 60]\n"
            "  content: {text: 第二页}\n",
            encoding="utf-8",
        )
        first = check_module.check_deck(self.root, page_spec="1")
        second = check_module.check_deck(self.root, page_spec="2")
        self.assertFalse(
            any("page-2.page" in item["location"] for item in first)
        )
        self.assertTrue(
            any("page-2.page" in item["location"] for item in second)
        )


class ConvertTests(unittest.TestCase):
    def test_convert_output_layout_and_media_rewrite(self):
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "red.png"
            Image.new("RGB", (12, 12), (200, 30, 30)).save(image)
            pptx = root / "deck.pptx"
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text_frame.text = "Hello"
            slide.shapes.add_picture(str(image), Inches(5), Inches(1), Inches(1), Inches(1))
            prs.save(pptx)

            out = root / "out"
            report = cli_module._convert_deck(pptx, out)
            self.assertEqual(report["pageCount"], 1)
            self.assertTrue((out / "deck.pptd").is_file())
            self.assertTrue((out / "pages" / "page-1.page").is_file())
            page = load_yaml(out / "pages" / "page-1.page")
            media_src = [e["src"] for e in page["elements"] if e.get("elementType") == "image"]
            self.assertTrue(media_src)
            self.assertTrue(all(src.startswith("media/") and (out / src).is_file() for src in media_src))
            self.assertNotIn("image-1", media_src[0])


class ScreenshotTests(unittest.TestCase):
    def test_screenshot_gated_on_renderer(self):
        if not cli_module._renderer_ready():
            self.skipTest("PPTD browser runtime or Chromium is unavailable")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _write_deck(
                root,
                "elements:\n- elementId: t\n  elementType: text\n  content: {text: hi}\n  bounds: [0, 0, 200, 60]\n",
            )
            out = root / "shots"
            entry = check_module._find_entry(root)
            written = cli_module._screenshot(entry, out, "1", 1.0, "png", None)
            self.assertEqual(written, ["page-1.png"])
            self.assertTrue((out / "pages" / "page-1.png").is_file())


if __name__ == "__main__":
    unittest.main()
