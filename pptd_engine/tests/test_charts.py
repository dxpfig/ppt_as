"""Observable OOXML/workbook/export behavior, not internal helper snapshots."""
import importlib
import io
import json
import unittest
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

def source(series=None, cols=None, rows=None, **extra):
    return dict(elementType='chart', elementId='chart-test', bounds=[20, 30, 600, 350], data={'cols': cols or ['x', 'a', 'b'], 'rows': rows or [['A', 2, 6], ['B', 4, 4]]}, series=series or [{'type': 'bar', 'encode': {'x': 'x', 'y': 'a'}}], **extra)

def export(element, **kwargs):
    module = importlib.import_module('pptd_engine.charts')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape, audit = module.add_chart(slide, element, {}, **kwargs)
    output = io.BytesIO()
    deck.save(output)
    return shape, audit, zipfile.ZipFile(io.BytesIO(output.getvalue()))

class ChartExportTests(unittest.TestCase):
    def test_label_layout_is_retained_in_editable_native_chart(self):
        _, _, package = export(source(xAxis={'label': {'rotation': 25, 'interval': 0}},
            dataLabels={'show': True, 'position': 'outsideEnd'}))
        chart = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(chart.xpath('//c:catAx/c:txPr/a:bodyPr/@rot', namespaces=NS), ['-1500000'])
        self.assertEqual(chart.xpath('//c:catAx/c:tickLblSkip/@val', namespaces=NS), ['1'])
        self.assertEqual(chart.xpath('//c:dLbls/c:dLblPos/@val', namespaces=NS), ['outEnd'])

    def test_native_chart_has_editable_data_and_original_source(self):
        element = source(rows=[['A', '2.5', 7], ['B', None, 4]])
        shape, audit, package = export(element)
        self.assertTrue(shape.has_chart)
        self.assertEqual(audit['mode'], 'native')
        chart = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(chart.xpath('//c:barChart/c:ser/c:val/c:numRef/c:numCache/c:pt/c:v/text()', namespaces=NS), ['2.5'])
        self.assertTrue(any(n.endswith('.xlsx') for n in package.namelist()))
        sources = [package.read(n) for n in package.namelist() if n.startswith('customXml/') and n.endswith('.xml')]
        self.assertTrue(any('"2.5"' in data.decode() and 'chart-test' in data.decode() for data in sources))

    def test_combo_and_secondary_value_axis_share_one_native_chart(self):
        series = [{'type': 'bar', 'encode': {'x': 'x', 'y': 'a'}}, {'type': 'line', 'encode': {'x': 'x', 'y': 'b'}, 'yAxisIndex': 1}]
        _, audit, package = export(source(series=series, yAxis=[{}, {'min': 0, 'max': 100}]))
        chart = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(audit['mode'], 'native')
        self.assertEqual(len(chart.xpath('//c:barChart', namespaces=NS)), 1)
        self.assertEqual(len(chart.xpath('//c:lineChart', namespaces=NS)), 1)
        self.assertEqual(len(chart.xpath('//c:valAx', namespaces=NS)), 2)
        self.assertIn('r', chart.xpath('//c:valAx/c:axPos/@val', namespaces=NS))
        self.assertIn('100', chart.xpath('//c:valAx/c:scaling/c:max/@val', namespaces=NS))

    def test_native_types_create_real_chart_plots(self):
        cases = [
            (source(series=[{'type': t, 'encode': {'x': 'x', 'y': 'a'}}]), tag)
            for t, tag in [('bar', 'barChart'), ('line', 'lineChart'), ('area', 'areaChart')]]
        cases += [
            (source(series=[{'type': 'scatter', 'encode': {'x': 'x', 'y': 'a'}}], rows=[[1, 2, 3], [2, 4, 6]]), 'scatterChart'),
            (source(series=[{'type': 'bubble', 'encode': {'x': 'x', 'y': 'a', 'size': 'b'}}], rows=[[1, 2, 3], [2, 4, 6]]), 'bubbleChart'),
            (source(series=[{'type': 'pie', 'encode': {'category': 'x', 'value': 'a'}, 'innerRadius': .4}]), 'doughnutChart'),
            (source(series=[{'type': 'radar', 'encode': {'category': 'x', 'y': 'a'}}], rows=[['A', 2, 6], ['B', 4, 4], ['C', 3, 5]]), 'radarChart'),
            (source(series=[{'type': 'candlestick', 'encode': {'x': 'x', 'high': 'h', 'low': 'l', 'close': 'c', 'open': 'o'}}], cols=['x','h','l','c','o'], rows=[['A', 9, 1, 7, 3], ['B', 9, 1, 3, 7]]), 'stockChart')]
        for element, tag in cases:
            with self.subTest(tag=tag):
                _, _, package = export(element)
                xml = etree.fromstring(package.read('ppt/charts/chart1.xml'))
                self.assertEqual(len(xml.xpath('//c:' + tag, namespaces=NS)), 1)

    def test_vector_export_requires_opt_in_and_has_editable_geometry_and_text(self):
        element = source(series=[{'type': 'sankey', 'encode': {'source': 'x', 'target': 'to', 'flow': 'a'}, 'dataLabels': {'show': True, 'content': 'category'}}], cols=['x','to','a'], rows=[['A', 'B', 2], ['A','C',3]])
        with self.assertRaisesRegex(ValueError, 'EditableVectorRequired'):
            export(element)
        shape, audit, package = export(element, allow_vector=True)
        self.assertEqual(audit['mode'], 'editable-vector')
        self.assertEqual(shape.shape_type, MSO_SHAPE_TYPE.GROUP)
        self.assertGreater(len(shape.shapes), 3)
        slide = package.read('ppt/slides/slide1.xml').decode()
        self.assertIn('custGeom', slide)
        self.assertNotIn('<p:pic>', slide)
        self.assertIn('<a:t>A</a:t>', slide)
        self.assertFalse(any(n.startswith('ppt/media/') for n in package.namelist()))

    def test_invalid_data_does_not_silently_become_zero(self):
        with self.assertRaisesRegex(ValueError, 'NonNumericValueError'):
            export(source(rows=[['A', 'wrong', 4]]))

    def test_percent_and_null_handling_reach_ooxml(self):
        element = source(series=[{'type': 'bar', 'encode': {'x': 'x', 'y': 'a'}, 'stack': 'percent'}, {'type': 'bar', 'encode': {'x': 'x', 'y': 'b'}, 'stack': 'percent'}])
        _, _, package = export(element)
        chart = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(chart.xpath('//c:barChart/c:grouping/@val', namespaces=NS), ['percentStacked'])
        _, _, package = export(source(series=[{'type': 'line', 'encode': {'x': 'x', 'y': 'a'}, 'nullHandling': 'connect'}]))
        self.assertIn(b'<c:dispBlanksAs val="span"', package.read('ppt/charts/chart1.xml'))

    def test_vector_clip_is_applied_to_editable_geometry(self):
        element = source(series=[{'type': 'sankey', 'encode': {'source': 'x', 'target': 'to', 'flow': 'a'}}], cols=['x','to','a'], rows=[['A','B',2]])
        svg = '<svg xmlns="http://www.w3.org/2000/svg" width="600" height="350"><defs><clipPath id="clip"><rect x="10" y="20" width="100" height="80"/></clipPath></defs><g clip-path="url(#clip)"><path d="M-20 0L200 0L200 150L-20 150Z" fill="#ff0000"/></g></svg>'
        shape, _, _ = export(element, allow_vector=True, svg=svg)
        colored = [s for s in shape.shapes if s._element.xpath('.//a:custGeom')]
        self.assertEqual(len(colored), 1)
        self.assertAlmostEqual(colored[0].left.pt, 10, places=3)
        self.assertAlmostEqual(colored[0].top.pt, 20, places=3)
        self.assertAlmostEqual(colored[0].width.pt, 100, places=3)
        self.assertAlmostEqual(colored[0].height.pt, 80, places=3)

    def test_all_five_special_types_export_editable_vectors(self):
        cases = [
            source(series=[{'type':'waterfall','encode':{'x':'x','y':'a'}}]),
            source(series=[{'type':'heatmap','encode':{'x':'x','y':'y','value':'v'}}], cols=['x','y','v'], rows=[['A','One',1],['B','Two',3]]),
        ]
        cases = cases[:2] + [source(series=[{'type':t,'encode':{'category':'x','value':'a'}}]) for t in ['treemap','sunburst']]
        cases.append(source(series=[{'type':'sankey','encode':{'source':'x','target':'to','flow':'a'}}],cols=['x','to','a'],rows=[['A','B',2]]))
        for element in cases:
            with self.subTest(type=element['series'][0]['type']):
                shape, audit, package = export(element, allow_vector=True)
                self.assertEqual(audit['mode'], 'editable-vector')
                self.assertGreater(len(shape.shapes), 1)
                self.assertNotIn(b'<p:pic>', package.read('ppt/slides/slide1.xml'))

    def test_native_image_frame_has_chart_local_image_relationship(self):
        from pathlib import Path
        element = source(fill={'type':'image','src':'bg.png','fit':{'mode':'contain'}}, shadow={'color':'#00000080','blur':4,'offset':[2,3]}, border={'color':'#123456','width':2})
        asset = Path(__file__).resolve().parents[2] / 'experiments/pptd-v2/fixture/media/fixture.png'
        _, _, package = export(element, asset_index={'bg.png':str(asset)})
        xml = package.read('ppt/charts/chart1.xml')
        self.assertIn(b'<a:blipFill>', xml)
        self.assertIn(b'<a:outerShdw', xml)
        self.assertIn(b'/image', package.read('ppt/charts/_rels/chart1.xml.rels'))

    def test_scatter_marker_inherits_series_border(self):
        element = source(series=[{'type':'scatter','encode':{'x':'x','y':'a'},'border':{'color':'#FF0000','width':3}}], rows=[[1,2,3],[2,4,6]])
        _, _, package = export(element)
        xml = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(xml.xpath('//c:marker/c:spPr/a:ln/@w', namespaces=NS), ['38100'])
        self.assertEqual(xml.xpath('//c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr/@val', namespaces=NS), ['FF0000'])

    def test_category_gap_is_ratio_of_category_slot_not_one_bar(self):
        element = source(series=[{'type':'bar','encode':{'x':'x','y':'a'}},{'type':'bar','encode':{'x':'x','y':'b'}}], categoryGap=.2, barGap=0)
        _, _, package = export(element)
        xml = etree.fromstring(package.read('ppt/charts/chart1.xml'))
        self.assertEqual(xml.xpath('//c:barChart/c:gapWidth/@val', namespaces=NS), ['50'])
        element['categoryGap'] = .99
        with self.assertRaisesRegex(ValueError, 'EditableVectorRequired'):
            export(element)

    def test_heatmap_colorbar_retains_gradient_inside_transformed_svg_clip(self):
        element = source(series=[{'type':'heatmap','encode':{'x':'x','y':'y','value':'v'}}], cols=['x','y','v'], rows=[['A','One',1],['B','Two',3]])
        _, _, package = export(element, allow_vector=True)
        xml = package.read('ppt/slides/slide1.xml')
        self.assertIn(b'<a:gradFill', xml)
        tree = etree.fromstring(xml)
        self.assertIn('16200000', tree.xpath('//a:gradFill/a:lin/@ang', namespaces=NS))

    def test_custom_svg_text_font_shorthand_survives_as_editable_text(self):
        element = source(series=[{'type':'sankey','encode':{'source':'x','target':'to','flow':'a'}}],cols=['x','to','a'],rows=[['A','B',2]])
        svg='<svg xmlns="http://www.w3.org/2000/svg" width="600" height="350"><text x="10" y="30" style="font: bold 19px DejaVu Sans" fill="#123456">Value</text></svg>'
        _, _, package = export(element,allow_vector=True,svg=svg)
        xml=etree.fromstring(package.read('ppt/slides/slide1.xml'))
        self.assertEqual(xml.xpath('//a:rPr/a:latin/@typeface',namespaces=NS),['DejaVu Sans'])
        self.assertEqual(xml.xpath('//a:rPr/@sz',namespaces=NS),['1900'])

if __name__ == '__main__':
    unittest.main()
