"""Behavioral package checks; PowerPoint rendering is a separate acceptance check."""
import copy
import io
import json
from pathlib import Path
import unittest
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
try:
    from pptd_engine.animations import apply_animations
except ImportError:
    apply_animations = None

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def validate_timing_structure(root):
    """Validate child order/cardinality and attributes against Microsoft SDK metadata.

    This is deliberately independent of the exporter. DrawingML color internals
    are outside this PresentationML schema slice; this is not a full OPC validator.
    """
    schema=json.loads((Path(__file__).parent/'animation_fixtures/presentationml_timing_schema.json').read_text())
    types={t['Name']:t for t in schema['types']}
    classes={t['ClassName']:t for t in schema['types']}
    def qualified(name):
        prefix,local=name.rsplit('/',1)[-1].split(':')
        return '{'+NS.get(prefix, 'urn:unsupported:'+prefix)+'}'+local
    def ancestors(desc):
        result=[desc]
        while desc.get('BaseClass') in classes:
            desc=classes[desc['BaseClass']];result.append(desc)
        return result
    def check(node,name):
        if name not in types: return
        desc=types[name];chain=ancestors(desc)
        attributes={a['QName'].lstrip(':'):a for t in chain for a in t.get('Attributes',[])}
        for attribute in node.attrib:
            if attribute not in attributes: raise AssertionError(f'{node.tag}: unrecognized attribute {attribute}')
        for key,attribute in attributes.items():
            if any(v.get('Name')=='RequiredValidator' for v in attribute.get('Validators',[])) and key not in node.attrib:
                raise AssertionError(f'{node.tag}: missing required attribute {key}')
        children=list(node)
        def once(rule,pos):
            if 'Name' in rule:
                return {pos+1} if pos<len(children) and children[pos].tag==qualified(rule['Name']) else set()
            if rule.get('Kind')=='Choice':
                return set().union(*(match(item,pos) for item in rule.get('Items',[])))
            positions={pos}
            for item in rule.get('Items',[]):
                positions=set().union(*(match(item,p) for p in positions))
            return positions
        def match(rule,pos):
            occurs=rule.get('Occurs')
            minimum,maximum=(1,1) if not occurs else (occurs[0].get('Min',0),occurs[0].get('Max',len(children)+1))
            accepted={pos} if minimum==0 else set();positions={pos}
            for count in range(1,maximum+1):
                positions=set().union(*(once(rule,p) for p in positions))
                if not positions: break
                if count>=minimum: accepted|=positions
            return accepted
        particle=next((t['Particle'] for t in chain if 'Particle' in t),None)
        if particle and len(children) not in match(particle,0):
            raise AssertionError(f'{node.tag}: child order/cardinality invalid: {[etree.QName(c).localname for c in children]}')
        candidates={qualified(c['Name']):c['Name'] for t in chain for c in t.get('Children',[])}
        for child in children:
            if child.tag not in candidates:
                raise AssertionError(f'{node.tag}: unexpected child {child.tag}')
            check(child,candidates[child.tag])
    check(root,'p:CT_SlideTiming/p:timing')

class AnimationExportTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(callable(apply_animations), 'Native PPTX animation exporter must exist')
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(5)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.a = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        self.b = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(1), Inches(1))
        self.shapes = {'a': self.a, 'b': self.b}

    def xml(self):
        out = io.BytesIO()
        self.prs.save(out)
        with zipfile.ZipFile(out) as z:
            return etree.fromstring(z.read('ppt/slides/slide1.xml'))

    def test_trigger_graph_delay_repeat_and_unique_shape_references(self):
        animations = [
            {'elementId':'a', 'effect':'fade-in', 'trigger':'withPrevious', 'delayMs':50, 'durationMs':100, 'repeat':2},
            {'elementId':'b', 'effect':'fade-in', 'trigger':'withPrevious', 'durationMs':200},
            {'elementId':'a', 'effect':'transparency', 'trigger':'afterPrevious', 'durationMs':100, 'amount':.3},
            {'elementId':'b', 'effect':'disappear'},
        ]
        original = copy.deepcopy(animations)
        audit = apply_animations(self.slide, animations, self.shapes, [1000,500])
        self.assertEqual(animations, original)
        self.assertEqual(audit['animationCount'], 4)
        tree = self.xml()
        groups = tree.xpath('//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par/p:cTn', namespaces=NS)
        self.assertEqual(len(groups),2)
        self.assertEqual(groups[0].xpath('p:stCondLst/p:cond/@delay',namespaces=NS),['0'])
        self.assertEqual(groups[1].xpath('p:stCondLst/p:cond/@delay',namespaces=NS),['indefinite'])
        effects = groups[0].xpath('p:childTnLst/p:par/p:cTn',namespaces=NS)
        self.assertEqual([n.xpath('p:stCondLst/p:cond/@delay',namespaces=NS)[0] for n in effects],['50','50','250'])
        self.assertEqual(effects[0].get('repeatCount'),'2000')
        self.assertEqual(effects[0].get('dur'),'100')
        ids=tree.xpath('//p:timing//p:cTn/@id',namespaces=NS)
        self.assertEqual(len(ids),len(set(ids)))
        self.assertEqual(set(tree.xpath('//p:timing//p:spTgt/@spid',namespaces=NS)),{str(self.a.shape_id),str(self.b.shape_id)})
        self.assertTrue(tree.xpath('//p:seq/p:nextCondLst/p:cond[@evt="onNext"]/p:tgtEl/p:sldTgt',namespaces=NS))

    def test_all_effects_emit_native_behavior_and_default_duration(self):
        cases = [
            ('appear',0,'set'), ('fade-in',500,'animEffect'), ('fly-in',500,'anim'), ('zoom-in',500,'animScale'),
            ('wipe-in',500,'animEffect'), ('float-in',500,'animEffect'), ('peek-in',500,'animEffect'), ('rise-in',1000,'anim'),
            ('pulse',600,'animScale'), ('grow-shrink',2000,'animScale'), ('spin',2000,'animRot'), ('teeter',1000,'animRot'),
            ('fill-color',2000,'animClr'), ('transparency',2000,'animEffect'), ('color-pulse',2000,'animClr'),
            ('disappear',0,'set'), ('fade-out',500,'animEffect'), ('fly-out',500,'anim'), ('zoom-out',500,'animScale'),
            ('wipe-out',500,'animEffect'), ('float-out',500,'animEffect'), ('motion-path',2000,'animMotion'),
        ]
        for effect,duration,behavior in cases:
            with self.subTest(effect=effect):
                audit=apply_animations(self.slide,[{'elementId':'a','effect':effect,'color':'#123456','amount':.4,'path':'M 0 0 L 200 -100'}],self.shapes,[1000,500])
                tree=self.xml()
                self.assertEqual(len(tree.xpath('p:timing',namespaces=NS)),1)
                self.assertTrue(tree.xpath(f'//p:timing//p:{behavior}',namespaces=NS))
                self.assertEqual(audit['animations'][0]['durationMs'],duration)
                self.assertFalse(tree.xpath('//p:timing//p:extLst',namespaces=NS))
                validate_timing_structure(tree.find('p:timing',NS))

    def test_color_opacity_easing_and_curve_coordinates_are_executable_values(self):
        apply_animations(self.slide,[
            {'elementId':'a','effect':'fill-color','color':'#123456','easing':'ease-in'},
            {'elementId':'a','effect':'color-pulse','color':'FF0000','trigger':'afterPrevious'},
            {'elementId':'a','effect':'transparency','amount':.3,'trigger':'afterPrevious'},
            {'elementId':'b','effect':'motion-path','path':'M 0 0 Q 100 -100 200 0 C 300 0 400 100 500 0 Z','easing':'ease-out'},
        ], self.shapes,[1000,500])
        tree=self.xml()
        self.assertEqual(tree.xpath('//p:animClr/p:to/a:srgbClr/@val',namespaces=NS),['123456','FF0000'])
        self.assertEqual(tree.xpath('//p:animClr[2]/p:cBhvr/p:cTn/@autoRev',namespaces=NS),[])  # different effect parents
        self.assertEqual(len(tree.xpath('//p:animClr/p:cBhvr/p:cTn[@autoRev="1"]',namespaces=NS)),1)
        self.assertTrue(tree.xpath('//p:set[p:cBhvr/p:attrNameLst/p:attrName="style.opacity"]/p:to/p:strVal[@val="0.3"]',namespaces=NS))
        self.assertTrue(tree.xpath('//p:animEffect[@filter="image"][@prLst="opacity: 0.3"]',namespaces=NS))
        self.assertTrue(tree.xpath('//p:cTn[@accel="100000"]',namespaces=NS))
        self.assertTrue(tree.xpath('//p:cTn[@decel="100000"]',namespaces=NS))
        path=tree.xpath('//p:animMotion/@path',namespaces=NS)[0]
        self.assertTrue(path.startswith('M 0 0 C '))  # quadratic converted to a cubic Office understands
        self.assertIn('0.2 0',path)
        self.assertIn('0.5 0',path)
        self.assertTrue(path.endswith('Z E'))
        self.assertNotIn('Q',path)

    def test_authored_rotation_color_baseline_and_pulse_restore_are_retained(self):
        from pptx.dml.color import RGBColor
        self.a.rotation = 30
        self.a.fill.solid()
        self.a.fill.fore_color.rgb = RGBColor.from_string('0000FF')
        apply_animations(self.slide,[
            {'elementId':'a','effect':'spin'},
            {'elementId':'a','effect':'fill-color','color':'FF0000','trigger':'afterPrevious'},
            {'elementId':'a','effect':'color-pulse','color':'00FF00','trigger':'afterPrevious'},
        ],self.shapes,[1000,500])
        tree=self.xml()
        self.assertEqual(tree.xpath('//p:animRot/@from',namespaces=NS),['1800000'])
        self.assertEqual(tree.xpath('//p:animClr/p:from/a:srgbClr/@val',namespaces=NS),['0000FF','FF0000'])
        pulse=tree.xpath('//p:animClr[p:cBhvr/p:cTn[@autoRev="1"]]',namespaces=NS)[0]
        self.assertEqual(pulse.xpath('p:cBhvr/p:cTn/@fill',namespaces=NS),['remove'])

    def test_exit_visibility_setter_is_not_truncated_by_its_time_container(self):
        apply_animations(self.slide,[{'elementId':'a','effect':'fly-out','durationMs':500}],self.shapes,[1000,500])
        tree=self.xml()
        effect=tree.xpath('//p:cTn[@presetClass="exit"]',namespaces=NS)[0]
        # Native sets require one millisecond of active time: they must finish
        # inside the declared effect so a subsequent entrance cannot be hidden.
        end_set=effect.xpath('p:childTnLst/p:set[p:to/p:strVal[@val="hidden"]]/p:cBhvr/p:cTn',namespaces=NS)[0]
        start=int(end_set.xpath('p:stCondLst/p:cond/@delay',namespaces=NS)[0])
        self.assertLessEqual(start+int(end_set.get('dur')),int(effect.get('dur','0')))

    def test_fill_of_nested_groups_targets_painted_descendants(self):
        inner=self.slide.shapes.add_group_shape([self.a,self.b])
        outer=self.slide.shapes.add_group_shape([inner])
        apply_animations(self.slide,[{'elementId':'outer','effect':'fill-color','color':'FF0000'}],{'outer':outer},[1000,500])
        tree=self.xml()
        self.assertEqual(set(tree.xpath('//p:animClr/p:cBhvr/p:tgtEl/p:spTgt/@spid',namespaces=NS)),{str(self.a.shape_id),str(self.b.shape_id)})

    def test_group_shapes_target_the_group_and_invalid_input_is_atomic(self):
        group=self.slide.shapes.add_group_shape([self.a,self.b])
        apply_animations(self.slide,[{'elementId':'group','effect':'pulse'}],{'group':group},[1000,500])
        self.assertEqual(set(self.xml().xpath('//p:timing//p:spTgt/@spid',namespaces=NS)),{str(group.shape_id)})
        before=etree.tostring(self.slide._element)
        for bad in [
            {'elementId':'missing','effect':'pulse'},
            {'elementId':'group','effect':'motion-path','path':'M 1 0 L 2 3'},
            {'elementId':'group','effect':'motion-path','path':'M 0 0 M 2 3'},
            {'elementId':'group','effect':'fill-color','color':'bad'},
            {'elementId':'group','effect':'float-in','direction':'left'},
        ]:
            with self.assertRaisesRegex(ValueError,r'animations\[0\]'):
                apply_animations(self.slide,[bad],{'group':group},[1000,500])
            self.assertEqual(etree.tostring(self.slide._element),before)

if __name__ == '__main__':
    unittest.main()
