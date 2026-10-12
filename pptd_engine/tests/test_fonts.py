import io
import struct
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from fontTools.ttLib import TTFont, TTCollection, newTable
from fontTools.ttLib.tables._f_v_a_r import Axis
from fontTools.ttLib.tables.TupleVariation import TupleVariation
from pptx import Presentation
from pptx.util import Pt
from lxml import etree
from pptd_engine.fonts import FontResolver, eot_encode, eot_decode, embed_fonts


class FontTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cache = Path(self.tmp.name)
        self.resolver = FontResolver(self.cache)
        self.font = self.resolver.resolve('Noto Sans')

    def tearDown(self):
        self.tmp.cleanup()

    def test_normalization_retains_timestamp_for_stable_shared_cache(self):
        original=TTFont(self.font.source.removeprefix('system:'))
        cached=TTFont(self.font.path)
        self.assertEqual(cached['head'].modified,original['head'].modified)
        second=FontResolver(self.cache).resolve('Noto Sans')
        self.assertEqual(second.path,self.font.path)

    def test_real_eot_roundtrip(self):
        raw = self.font.path.read_bytes()
        eot = eot_encode(raw)
        self.assertEqual(struct.unpack_from('<I', eot)[0], len(eot))
        self.assertEqual(struct.unpack_from('<H', eot, 34)[0], 0x504c)
        self.assertEqual(eot_decode(eot), raw)
        self.assertEqual(TTFont(io.BytesIO(eot_decode(eot))).getGlyphOrder(), TTFont(io.BytesIO(raw)).getGlyphOrder())

    def test_restricted_permissions_rejected(self):
        for flag in (2, 4, 0x200):
            font = TTFont(self.font.path)
            font['OS/2'].fsType = flag
            buf = io.BytesIO(); font.save(buf)
            with self.assertRaisesRegex(ValueError, 'embedding|editable'):
                eot_encode(buf.getvalue())

    def test_full_font_allowed_without_subsetting(self):
        font = TTFont(self.font.path); font['OS/2'].fsType = 0x108
        buf = io.BytesIO(); font.save(buf)
        self.assertEqual(eot_decode(eot_encode(buf.getvalue())), buf.getvalue())

    def test_no_silent_font_substitution_or_missing_glyph(self):
        with self.assertRaisesRegex(ValueError, 'not installed'):
            self.resolver.resolve('Definitely Missing Font 987')
        with self.assertRaisesRegex(ValueError, 'glyph'):
            self.font.check_coverage('测试')
        self.font.check_coverage('Editable 123')

    def test_embedded_faces_follow_schema_order_independent_of_usage_order(self):
        bold=self.resolver.resolve('Noto Sans',700)
        italic=self.resolver.resolve('Noto Sans',400,True)
        dest=self.cache/'ordered.pptx';Presentation().save(dest)
        embed_fonts(dest,[italic,bold,self.font])
        with ZipFile(dest) as z:
            root=etree.fromstring(z.read('ppt/presentation.xml'))
        faces=root.xpath('//*[local-name()="embeddedFont"]/*')
        self.assertEqual([etree.QName(face).localname for face in faces],['font','regular','bold','italic'])

    def test_reembedding_replaces_existing_slot_and_does_not_leave_orphan_font_parts(self):
        dest=self.cache/'repeat.pptx';Presentation().save(dest)
        embed_fonts(dest,[self.font]);embed_fonts(dest,[self.font])
        with ZipFile(dest) as z:
            root=etree.fromstring(z.read('ppt/presentation.xml'))
            rels=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
            faces=root.xpath('//*[local-name()="embeddedFont"]/*[local-name()="regular"]')
            font_rels=[rel for rel in rels if rel.get('Type').endswith('/font')]
            self.assertEqual(len(faces),1)
            self.assertEqual(len(font_rels),1)
            self.assertEqual(len([part for part in z.namelist() if part.startswith('ppt/fonts/')]),1)
            rid=faces[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            self.assertEqual(font_rels[0].get('Id'),rid)
            self.assertEqual(eot_decode(z.read('ppt/'+font_rels[0].get('Target'))),self.font.path.read_bytes())

    def test_custom_collection_selects_matching_face_instead_of_first_font(self):
        bold=self.resolver.resolve('Noto Sans',700)
        collection=TTCollection();collection.fonts=[TTFont(self.font.path),TTFont(bold.path)]
        buffer=io.BytesIO();collection.save(buffer)
        resolver=FontResolver(self.cache/'ttc',[{'family':'Noto Sans','src':'https://example.test/fonts.ttc'}])
        resolver._download=lambda url:buffer.getvalue()
        face=resolver.resolve('Noto Sans',700)
        self.assertEqual(TTFont(face.path)['OS/2'].usWeightClass,700)
        self.assertEqual(face.slot,'bold')

    def test_variable_italic_axis_changes_outline_and_sets_static_face_flags(self):
        font=TTFont(self.font.path)
        font['fvar']=newTable('fvar');font['fvar'].axes=[];font['fvar'].instances=[]
        axis=Axis();axis.axisTag='ital';axis.minValue=0;axis.defaultValue=0;axis.maxValue=1;axis.flags=0;axis.axisNameID=256
        font['name'].setName('Italic axis',256,3,1,0x409);font['fvar'].axes.append(axis)
        coordinates=font['glyf']['A'].getCoordinates(font['glyf'])[0]
        original_x=coordinates[0][0]
        font['gvar']=newTable('gvar');font['gvar'].variations={glyph:[] for glyph in font.getGlyphOrder()}
        deltas=[(20,0)]+[(0,0)]*(len(coordinates)+3)
        font['gvar'].variations['A']=[TupleVariation({'ital':(0,1,1)},deltas)]
        buffer=io.BytesIO();font.save(buffer)
        resolver=FontResolver(self.cache/'variable',[{'family':'Noto Sans','src':'https://example.test/variable.ttf'}])
        resolver._download=lambda url:buffer.getvalue()
        face=resolver.resolve('Noto Sans',400,True)
        static=TTFont(face.path)
        self.assertNotIn('fvar',static)
        self.assertTrue(static['OS/2'].fsSelection&1)
        self.assertEqual(static['glyf']['A'].getCoordinates(static['glyf'])[0][0][0],original_x+20)
        self.assertEqual(face.slot,'italic')

    def test_stix_math_resource_must_contain_math_layout_table(self):
        font=TTFont(self.font.path)
        for record in font['name'].names:
            if record.nameID in (1,16):record.string='STIX Two Math'.encode(record.getEncoding())
        buffer=io.BytesIO();font.save(buffer)
        resolver=FontResolver(self.cache/'math',[{'family':'STIX Two Math','src':'https://example.test/fake-math.ttf'}])
        resolver._download=lambda url:buffer.getvalue()
        with self.assertRaisesRegex(ValueError,'MATH'):
            resolver.resolve('STIX Two Math')

    def test_package_has_correct_parts_and_ea_font(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        r = slide.shapes.add_textbox(Pt(10),Pt(10),Pt(200),Pt(30)).text_frame.paragraphs[0].add_run()
        r.text = 'Font 123'; r.font.name = 'Noto Sans'
        dest = self.cache/'test.pptx'; prs.save(dest)
        audit = embed_fonts(dest, [self.font])
        self.assertEqual(audit[0]['format'], 'EOT-uncompressed')
        with ZipFile(dest) as z:
            ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            xml=etree.fromstring(z.read('ppt/presentation.xml'))
            self.assertEqual(xml.get('embedTrueTypeFonts'), '1')
            self.assertEqual(xml.xpath('p:embeddedFontLst/p:embeddedFont/p:font/@typeface',namespaces=ns), ['Noto Sans'])
            rels=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
            rel=[r for r in rels if r.get('Type').endswith('/font')][0]
            self.assertEqual(eot_decode(z.read('ppt/'+rel.get('Target'))), self.font.path.read_bytes())
            self.assertIn(b'application/x-fontdata', z.read('[Content_Types].xml'))
        self.assertEqual(len(Presentation(dest).slides),1)

if __name__ == '__main__': unittest.main()
