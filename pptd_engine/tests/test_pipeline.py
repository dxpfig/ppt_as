"""Cross-module acceptance: authoritative PPTD -> native/editable Office objects."""
from pathlib import Path
import tempfile
import unittest
import hashlib
from zipfile import ZipFile
from lxml import etree
from pptd_engine.pptx import Exporter
from pptd_engine.source_archive import recover_source
from pptd_engine.fonts import eot_decode
from pptd_engine.document import load_yaml

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'experiments/pptd-v2/advanced-fixture/advanced.pptd'

def hashes(folder):return {str(p.relative_to(folder)):hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.rglob('*') if p.is_file()}

class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name)
        cls.before=hashes(SOURCE.parent)
        cls.report=Exporter(SOURCE,cls.out/'export',allow_vector=True).run()
        cls.file=cls.out/'export/export.pptx'

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_all_objects_and_contracts_survive(self):
        report=self.report
        self.assertEqual(report['unsupported'],[])
        self.assertEqual(report['summary'],{'native-object':80,'editable-vector':6})
        self.assertEqual(len(report['ooxmlEvidence']['slides']),18)
        self.assertEqual(sum(x['nativeEquations'] for x in report['ooxmlEvidence']['slides']),6)
        self.assertEqual(sum(x['animationCount'] for x in report['animations']),22)
        effects={a['effect'] for p in report['animations'] for a in p['animations']}
        self.assertEqual(len(effects),22)
        self.assertEqual(len(report['charts']),15)
        self.assertEqual(len({t for c in report['charts'] for t in c['types']}),13)
        self.assertEqual(sum(c['mode']=='native' for c in report['charts']),9)
        self.assertEqual(len(report['ooxmlEvidence']['embeddedWorkbooks']),9)
        self.assertEqual(sum(s['nativePictures'] for s in report['ooxmlEvidence']['slides']),0)

    def test_actual_fonts_and_full_source_recovery(self):
        self.assertEqual(hashes(SOURCE.parent),self.before)
        recovered=recover_source(self.file,self.out/'recovered')
        self.assertEqual(recovered.name,SOURCE.name)
        self.assertEqual(hashes(recovered.parent),self.before)
        with ZipFile(self.file) as z:
            for font in self.report['fonts']:
                sfnt=eot_decode(z.read(font['part']))
                self.assertEqual(hashlib.sha256(sfnt).hexdigest(),font['sfntSha256'])
            self.assertIn('JetBrains Mono',[f['family'] for f in self.report['fonts']])
            self.assertIn('Noto Sans CJK SC',[f['family'] for f in self.report['fonts']])
            relations=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
            for relation in relations:
                if relation.get('Type').endswith('/font'):
                    self.assertIn('ppt/'+relation.get('Target'),z.namelist())

    def test_bad_formula_is_a_located_failure(self):
        from pptd_engine.pptx import Report
        from pptx import Presentation
        exporter=Exporter(SOURCE,self.out/'bad')
        slide=exporter.prs.slides.add_slide(exporter.prs.slide_layouts[6])
        with self.assertRaisesRegex(ValueError,'Invalid math'):
            exporter.element(slide,{'elementId':'bad','elementType':'text','bounds':[0,0,100,30],'content':{'text':r'\(\unknown{x}\)'}},'pages/01.page#bad')

    def test_vector_policy_is_explicit(self):
        from pptd_engine.charts import add_chart
        exporter=Exporter(SOURCE,self.out/'native-only')
        slide=exporter.prs.slides.add_slide(exporter.prs.slide_layouts[6])
        page=load_yaml(SOURCE.parent/'pages/14.page')
        with self.assertRaisesRegex(ValueError,'[Vv]ector|allow_vector'):
            add_chart(slide,page['elements'][1],exporter.theme)

if __name__=='__main__':unittest.main()
