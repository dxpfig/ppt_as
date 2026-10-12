"""Native HEX8 colors must retain alpha in editable DrawingML objects."""
import io
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

from pptd_engine.pptx import Exporter, NS


class NativeColorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "document.pptd"
        self.source.write_text(json.dumps({"version": "v2", "size": [960, 540], "pages": ["one.page"],
                                          "theme": {"colors": {"clear": "#00000000", "half": "#12345680", "alias": "$half"}}}), encoding="utf-8")
        self.exporter = Exporter(self.source, self.root / "output")
        self.slide = self.exporter.prs.slides.add_slide(self.exporter.prs.slide_layouts[6])

    def tearDown(self):
        self.tmp.cleanup()

    def test_transparent_fills_and_partial_borders_survive_native_pptx_save(self):
        shape = self.exporter.element(self.slide, {
            "elementId": "transparent-card", "elementType": "shape", "shapeName": "rect", "bounds": [50, 50, 300, 200],
            "fill": {"type": "solid", "color": "$clear"},
            "border": {"color": "$alias", "width": 0, "style": "solid"},
        }, "one.page#transparent-card")
        self.assertEqual(shape._element.xpath("./p:spPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["0"])
        self.assertEqual(shape._element.xpath("./p:spPr/a:ln/a:solidFill/a:srgbClr/a:alpha/@val"), ["50196"])
        self.assertEqual(shape._element.xpath("./p:spPr/a:ln/@w"), ["0"])
        stream = io.BytesIO(); self.exporter.prs.save(stream)
        with ZipFile(io.BytesIO(stream.getvalue())) as archive:
            xml = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
            self.assertEqual(xml.xpath("//p:sp/p:spPr/a:solidFill/a:srgbClr[@val='000000']/a:alpha/@val", namespaces=NS), ["0"])
            self.assertEqual(xml.xpath("//p:sp/p:spPr/a:ln/a:solidFill/a:srgbClr[@val='123456']/a:alpha/@val", namespaces=NS), ["50196"])
            self.assertEqual(len(xml.xpath("//p:pic", namespaces=NS)), 0)
        reopened = Presentation(io.BytesIO(stream.getvalue())).slides[0].shapes[0]
        self.assertEqual(reopened._element.xpath("./p:spPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["0"])
        self.assertEqual(self.exporter.report.data["unsupported"], [])

    def test_text_table_cells_and_slide_background_keep_native_alpha(self):
        self.exporter.fill(self.slide.background.fill, {"type": "solid", "color": "#55667740"}, "one.page")
        text = self.exporter.element(self.slide, {"elementId": "caption", "elementType": "text", "bounds": [10, 10, 400, 60],
                                                  "content": {"text": "Translucent text", "color": "$half"}}, "one.page#caption")
        table = self.exporter.element(self.slide, {"elementId": "table", "elementType": "table", "bounds": [10, 120, 400, 100],
                                                   "columnWidths": [1], "rowHeights": [1], "rows": [[{"text": "Cell", "fill": {"type": "solid", "color": "$half"}, "color": "#11223340"}]]}, "one.page#table")
        self.assertEqual(self.slide._element.xpath("./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["25098"])
        self.assertEqual(text._element.xpath(".//a:r/a:rPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["50196"])
        self.assertEqual(table._element.xpath(".//a:tcPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["50196"])
        self.assertEqual(table._element.xpath(".//a:r/a:rPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["25098"])
        self.assertEqual(self.exporter.report.data["unsupported"], [])

    def test_replacing_translucent_color_does_not_leave_stale_alpha(self):
        fill = self.slide.background.fill
        self.exporter.fill(fill, {"type": "solid", "color": "#12345680"}, "one.page")
        self.assertEqual(self.slide._element.xpath(".//p:bg//a:alpha/@val"), ["50196"])
        for opaque in ["#123456", "#123456FF"]:
            self.exporter.fill(fill, {"type": "solid", "color": opaque}, "one.page")
            self.assertEqual(self.slide._element.xpath(".//p:bg//a:alpha/@val"), [])
        self.exporter.fill(fill, {"type": "solid", "color": "#12345600"}, "one.page")
        self.assertEqual(self.slide._element.xpath(".//p:bg//a:alpha/@val"), ["0"])


if __name__ == "__main__":
    unittest.main()
